"""Safe, typed Excel export with separate audit and provenance sheets."""

from __future__ import annotations

import csv
import math
import re
import sys
from pathlib import Path
from typing import Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.cell_style import StyleArray

from .models import Result
from .values import date_value, is_identifier, number

NAVY = "17324D"
YELLOW = PatternFill("solid", fgColor="FFF0B3")
EXCEL_TEXT_LIMIT = 32767
XML_ESCAPE = re.compile(r"_[xX][0-9a-fA-F]{4}_")


def excel_text(value: str) -> str:
    """Encode forbidden XML characters without confusing literal escape sequences."""
    return "".join(
        "\\\\"
        if character == "\\"
        else "\\u005f"
        if character == "_" and XML_ESCAPE.match(value, index)
        else character
        if character in "\t\n\r"
        or 0x20 <= ord(character) <= 0xD7FF
        or 0xE000 <= ord(character) <= 0xFFFD
        or 0x10000 <= ord(character) <= 0x10FFFF
        else f"\\u{ord(character):04x}"
        for index, character in enumerate(value)
    )


def _set_text(cell, encoded: str) -> None:
    """Avoid openpyxl's silent string truncation and formula interpretation."""
    if len(encoded) > EXCEL_TEXT_LIMIT:
        raise ValueError("Excel text exceeds one cell; use the overflow-preserving exporter.")
    cell.value = encoded
    cell.data_type = "s"


def safe_text(cell, value: str) -> None:
    """Set an explicit string type, preventing spreadsheet formula injection."""
    _set_text(cell, excel_text(value))


class _TextWriter:
    """Preserve oversized values in numbered chunks rather than silently cutting them."""

    def __init__(self) -> None:
        self.overflow: list[tuple[str, str, str, str]] = []

    def write(self, cell, value: str) -> None:
        encoded = excel_text(value)
        if len(encoded) > EXCEL_TEXT_LIMIT:
            identifier = f"text-{len(self.overflow) + 1}"
            self.overflow.append((identifier, cell.parent.title, cell.coordinate, encoded))
            marker = f"[Full text: Long Text / {identifier}; this cell is a preview] "
            encoded = marker + encoded[: EXCEL_TEXT_LIMIT - len(marker)]
        _set_text(cell, encoded)

    def append(self, sheet, values) -> None:
        # Worksheet.append validates strings immediately, before a later cleanup can run.
        sheet.append([None if isinstance(value, str) else value for value in values])
        # Pinned openpyxl tracks the appended row here; max_row rescans every existing cell.
        row = sheet._current_row
        for column, value in enumerate(values, 1):
            if isinstance(value, str):
                self.write(sheet.cell(row, column), value)

    def finish(self, workbook) -> None:
        if not self.overflow:
            return
        sheet = workbook.create_sheet("Long Text")
        self.append(
            sheet, ["Record", "Source sheet", "Source cell", "Chunk", "Chunks", "Encoded text"]
        )
        for identifier, source_sheet, source_cell, encoded in self.overflow:
            chunks = [encoded[start : start + 30000] for start in range(0, len(encoded), 30000)]
            for index, chunk in enumerate(chunks, 1):
                sheet.append([identifier, source_sheet, source_cell, index, len(chunks), None])
                # Chunks are already encoded; encoding again would change the original value.
                _set_text(sheet.cell(sheet._current_row, 6), chunk)


def write_value(cell, value: str, text_writer: Callable = safe_text) -> None:
    """Type only unambiguous values; preserve IDs and precision-sensitive data."""
    if is_identifier(value):
        text_writer(cell, value)
        cell.number_format = "@"
    elif (parsed := number(value)) is not None and len(parsed.as_tuple().digits) <= 15:
        numeric = float(parsed)
        # Excel does not retain subnormal doubles. Do not turn a nonzero reading into zero.
        if not math.isfinite(numeric) or (parsed and abs(numeric) < sys.float_info.min):
            text_writer(cell, value)
            return
        decimal_places = max(0, -parsed.as_tuple().exponent)
        if value.strip().endswith("%"):
            precision = max(2 if "." in value else 0, decimal_places - 2)
            if precision > 12:
                text_writer(cell, value)
                return
            cell.number_format = "0" + ("." + "0" * precision if precision else "") + "%"
        else:
            precision = (
                min(12, len(value.split(".", 1)[1].split()[0].rstrip(")"))) if "." in value else 0
            )
            cell.number_format = (
                "General"
                if decimal_places > 12
                else "#,##0" + ("." + "0" * precision if precision else "")
            )
        cell.value = numeric if parsed % 1 else int(parsed)
    elif (parsed_date := date_value(value)) is not None:
        cell.value = parsed_date
        cell.number_format = "yyyy-mm-dd"
    else:
        text_writer(cell, value)


def style_sheet(sheet) -> None:
    """Apply readable, compact formatting to every populated sheet."""
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for column in sheet.columns:
        length = max(len(str(c.value or "")) for c in column)
        sheet.column_dimensions[column[0].column_letter].width = min(65, max(13, length + 2))
    # Register each shared style once. Assigning Font/Alignment through descriptors for
    # every evidence cell repeatedly hashes and compares the same style definitions.
    font_id = sheet.parent._fonts.add(Font(name="Calibri", size=11))
    alignment_id = sheet.parent._alignments.add(Alignment(vertical="top", wrap_text=True))
    for row in sheet.iter_rows(min_row=2):
        line_count = 1
        for cell in row:
            if cell._style is None:
                cell._style = StyleArray()
            cell._style.fontId = font_id
            cell._style.alignmentId = alignment_id
            width = sheet.column_dimensions[cell.column_letter].width
            line_count = max(line_count, math.ceil(len(str(cell.value or "")) / max(1, width - 2)))
        sheet.row_dimensions[row[0].row].height = min(409, 17 * line_count)


def export_workbook(result: Result, output: str | Path, *, reviewed: bool = False) -> Path:
    """Write a draft or quality-gated reviewed workbook with complete evidence accounting."""
    from .review import cell_evidence, quality_state, review_current

    if reviewed and (not review_current(result) or quality_state(result)["status"] != "reviewed"):
        raise ValueError(
            "Reviewed Excel is available only after every actionable item is checked "
            "and the current snapshot has a source-check attestation."
        )
    workbook = Workbook()
    workbook.remove(workbook.active)
    strings = _TextWriter()
    provenance = []
    evidence_rows = []

    def append_cell_evidence(
        sheet,
        excel_row: int,
        excel_column: int,
        table_index: int,
        source_row: int,
        source_column: int,
    ) -> None:
        evidence = cell_evidence(result, table_index, source_row, source_column)
        if not evidence["value"].strip():
            return
        sheet.cell(
            excel_row, excel_column
        ).hyperlink = f"#'Cell Evidence'!A{len(evidence_rows) + 2}"
        alternatives = " | ".join(
            f"{item['kind']}: {item['text']} ({item['engine']})"
            for item in evidence["alternative_readings"]
        )
        records = evidence["sources"] or [
            {
                "page": evidence["location"]["page"],
                "bbox": None,
                "text": evidence["original_reading"],
                "engine": None,
                "confidence": None,
                "role": "unlocated",
                "confidence_evidence": evidence["confidence"],
            }
        ]
        for record in records:
            confidence = record["confidence_evidence"]
            evidence_rows.append(
                [
                    evidence["evidence_id"],
                    sheet.title,
                    excel_row,
                    excel_column,
                    evidence["value"],
                    evidence["original_value"],
                    evidence["original_reading"],
                    record.get("page") or evidence["location"]["page"],
                    *(record.get("bbox") or [None] * 4),
                    record.get("engine"),
                    confidence["kind"],
                    confidence["score"],
                    confidence["meaning"],
                    record.get("text"),
                    record.get("initial_reading", ""),
                    record.get("role", "extracted"),
                    evidence["location"]["evidence_state"],
                    evidence["location"]["status"],
                    evidence["location"]["reason"],
                    evidence["correction_count"],
                    alternatives,
                ]
            )

    for index, table in enumerate(result.tables, 1):
        sheet = workbook.create_sheet(f"Data {index}")
        for col, header in enumerate(table.headers, 1):
            strings.write(sheet.cell(1, col), header)
            append_cell_evidence(sheet, 1, col, index - 1, -1, col - 1)
        for row_no, row in enumerate(table.rows, 2):
            for col, value in enumerate(row, 1):
                write_value(sheet.cell(row_no, col), value, strings.write)
                append_cell_evidence(sheet, row_no, col, index - 1, row_no - 2, col - 1)
            for page, source_row in table.provenance[row_no - 2]:
                provenance.append(
                    [
                        sheet.title,
                        row_no,
                        page,
                        source_row,
                        table.row_engines[row_no - 2],
                    ]
                )
        for issue in result.issues:
            if issue.table != index:
                continue
            for r, sources in enumerate(table.provenance, 2):
                if issue.row is not None and (issue.page, issue.row) not in sources:
                    continue
                if issue.row is None and not any(p == issue.page for p, _ in sources):
                    continue
                columns = [issue.column] if issue.column else range(1, len(table.headers) + 1)
                for col in columns:
                    sheet.cell(r, col).fill = YELLOW
    sheet = workbook.create_sheet("Validation Report")
    strings.append(
        sheet,
        [
            "Severity",
            "Page",
            "Source row",
            "Column",
            "Table",
            "Code",
            "Reason",
            "Suggested action",
        ],
    )
    for issue in result.issues:
        strings.append(
            sheet,
            [
                issue.severity,
                issue.page,
                issue.row,
                issue.column,
                issue.table,
                issue.code,
                issue.reason,
                issue.suggestion,
            ],
        )
    if not result.issues:
        strings.append(
            sheet,
            [
                "info",
                None,
                None,
                None,
                None,
                "no_findings",
                "No implemented rule raised a finding. This is not proof of accuracy.",
                "Compare with the source.",
            ],
        )
    sheet = workbook.create_sheet("Page Audit")
    strings.append(
        sheet,
        [
            "Page",
            "Route",
            "Engine",
            "Confidence",
            "Seconds",
            "Strategy",
            "Candidate ID",
            "Phase",
            "Execution",
            "Rank",
            "Selection reason",
            "Selected",
            "Tables",
            "Cells",
            "Findings",
            "Score",
            "Error",
        ],
    )
    for page in result.pages:
        if not page.attempts:
            strings.append(
                sheet,
                [
                    page.page,
                    page.route,
                    page.engine,
                    page.confidence,
                    page.seconds,
                    "",
                    "",
                    "",
                    "",
                    None,
                    "",
                    True,
                    0,
                    0,
                    0,
                    0,
                    page.error,
                ],
            )
        for attempt in page.attempts:
            strings.append(
                sheet,
                [
                    page.page,
                    page.route,
                    attempt.engine,
                    page.confidence,
                    round(attempt.seconds, 3),
                    attempt.strategy,
                    attempt.candidate_id,
                    attempt.phase,
                    attempt.execution,
                    attempt.rank,
                    attempt.selection_reason,
                    attempt.selected,
                    attempt.tables,
                    attempt.cells,
                    attempt.issues,
                    attempt.score,
                    attempt.error,
                ],
            )
    sheet = workbook.create_sheet("Join Receipt")
    strings.append(
        sheet,
        [
            "Data sheet",
            "Source pages",
            "Page joins",
            "Repeated headers removed",
            "Split rows recovered",
            "PDF page",
            "Source rows on page",
            "Output rows",
        ],
    )
    for index, table in enumerate(result.tables, 1):
        receipt = table.join_receipt
        pages = receipt.get("pages") or sorted(
            {page for sources in table.provenance for page, _ in sources} or {table.page}
        )
        rows_by_page = receipt.get("rows_by_page") or [
            {
                "page": page,
                "rows": sum(
                    source_page == page
                    for sources in table.provenance
                    for source_page, _ in sources
                ),
            }
            for page in pages
        ]
        for page_rows in rows_by_page:
            strings.append(
                sheet,
                [
                    f"Data {index}",
                    ", ".join(str(page) for page in pages),
                    receipt.get("page_joins", max(0, len(pages) - 1)),
                    receipt.get("repeated_headers_removed", 0),
                    receipt.get("split_rows_recovered", 0),
                    page_rows["page"],
                    page_rows["rows"],
                    receipt.get("output_rows", len(table.rows)),
                ],
            )
    sheet = workbook.create_sheet("Provenance")
    strings.append(
        sheet, ["Data sheet", "Excel row", "PDF page", "Source table data row", "Engine"]
    )
    for row in provenance:
        strings.append(sheet, row)
    sheet = workbook.create_sheet("Read me")
    strings.append(sheet, ["Field", "Value"])
    strings.append(sheet, ["Source", result.source])
    strings.append(sheet, ["Runtime seconds", round(result.seconds, 3)])
    strings.append(
        sheet,
        [
            "Text encoding",
            "All XLSX string fields use a reversible encoding: a literal backslash is doubled; "
            "XML-forbidden code points use a backslash followed by u and four hexadecimal digits. "
            "An underscore starting an OOXML escape-like sequence uses the same Unicode encoding. "
            "Decode left to right, treating each doubled backslash as one literal backslash. "
            "The full audit JSON and saved workspace retain the original strings.",
        ],
    )
    strings.append(
        sheet,
        [
            "Long text",
            "Values exceeding Excel's 32767-character cell limit show an explicit preview. "
            "The Long Text sheet preserves the complete encoded value in numbered chunks. "
            "Join its chunks in order before decoding the text encoding.",
        ],
    )
    strings.append(
        sheet,
        [
            "Small numbers",
            "Numbers below Excel's normal floating-point range remain exact text. "
            "Numbers needing more than 12 decimal places use General formatting; "
            "percentages needing more than 12 display decimals remain text, avoiding a false zero.",
        ],
    )
    strings.append(
        sheet,
        [
            "Yellow cells",
            "Review required. See Validation Report for the reason and source row.",
        ],
    )
    strings.append(
        sheet,
        [
            "Source row",
            "Data row within the extracted source table, excluding detected headers.",
        ],
    )
    strings.append(
        sheet,
        [
            "Confidence",
            "Extraction heuristic / OCR score, not a calibrated probability of correctness.",
        ],
    )
    strings.append(
        sheet,
        [
            "Dates",
            "ISO and unambiguous day-first dates are typed. Ambiguous dates stay text.",
        ],
    )
    strings.append(
        sheet,
        [
            "Safety",
            "Source strings are stored as text. No source formulas are executed.",
        ],
    )
    strings.append(
        sheet,
        [
            "Cell evidence",
            "Ctrl+click a linked data cell for its source box, confidence meaning, original readings, alternatives and correction count.",
        ],
    )
    strings.append(
        sheet,
        [
            "Repair proposals",
            "Arithmetic-derived suggestions only. Source documents may contain errors. Never applied automatically.",
        ],
    )
    sheet = workbook.create_sheet("Cell Evidence")
    strings.append(
        sheet,
        [
            "Evidence ID",
            "Data sheet",
            "Excel row",
            "Column",
            "Working value",
            "Original extracted value",
            "Original source reading",
            "Source page",
            "x0",
            "y0",
            "x1",
            "y1",
            "Engine",
            "Confidence type",
            "Recognition score",
            "Confidence meaning",
            "Reading",
            "Initial reading",
            "Role",
            "Evidence state",
            "Location status",
            "Location reason",
            "Corrections",
            "Alternative readings",
        ],
    )
    for row in evidence_rows:
        strings.append(sheet, row)
    from .constraints import coverage
    from .review import evidence_coverage, fingerprint

    sheet = workbook.create_sheet("Verification")
    strings.append(sheet, ["Evidence", "Value"])
    strings.append(sheet, ["Source SHA-256", result.source_sha256])
    strings.append(sheet, ["Extraction snapshot SHA-256", fingerprint(result)])
    strings.append(
        sheet,
        [
            "Review status",
            "Reviewed current snapshot" if review_current(result) else "Unreviewed",
        ],
    )
    strings.append(
        sheet,
        [
            "Review meaning",
            "A human source-check attestation; not independent accuracy certification.",
        ],
    )
    strings.append(
        sheet, ["Review note", result.review.get("note", "") if review_current(result) else ""]
    )
    strings.append(sheet, ["Export type", "Reviewed Excel" if reviewed else "Draft Excel"])
    evidence = evidence_coverage(result)
    strings.append(sheet, ["Non-empty exported cells", evidence["nonempty_cells"]])
    strings.append(sheet, ["Cells with source boxes", evidence["located"]])
    strings.append(sheet, ["Cells requiring full-page evidence review", evidence["unlocated"]])
    strings.append(sheet, ["Cell source-box coverage", evidence["coverage"]])
    sheet.cell(sheet.max_row, 2).number_format = "0.0%"
    for name, value in coverage(result).items():
        strings.append(sheet, [name.replace("_", " "), value])
    sheet = workbook.create_sheet("Constraints")
    strings.append(
        sheet,
        [
            "ID",
            "Table",
            "Inferred relation",
            "Checked rows",
            "Matching rows",
            "Empirical agreement",
            "Violations",
            "Meaning",
        ],
    )
    for relation in result.constraints:
        strings.append(
            sheet,
            [
                relation["id"],
                relation["table"],
                relation["formula"],
                relation["checked"],
                relation["matched"],
                relation["support"],
                len(relation["violations"]),
                "Inferred from this document; not a probability of correctness.",
            ],
        )
        sheet.cell(sheet.max_row, 6).number_format = "0.0%"
    sheet = workbook.create_sheet("Repair Proposals")
    strings.append(
        sheet,
        [
            "Table",
            "Excel row",
            "Column",
            "Observed",
            "Proposed",
            "Supporting equations",
            "Evidence",
            "State",
            "Caution",
        ],
    )
    for proposal in result.proposals:
        strings.append(
            sheet,
            [
                proposal["table"],
                proposal["row"] + 2,
                proposal["column"] + 1,
                proposal["observed"],
                proposal["proposed"],
                proposal["support_count"],
                "; ".join(proposal["supporting_equations"]),
                proposal["status"],
                proposal["basis"],
            ],
        )
    sheet = workbook.create_sheet("Correction History")
    strings.append(
        sheet,
        [
            "Evidence ID",
            "Time UTC",
            "Table",
            "Excel row",
            "Column",
            "Before",
            "After",
            "Reason",
            "Previous hash",
            "Event hash",
        ],
    )
    for event in result.history:
        strings.append(
            sheet,
            [
                (
                    f"T{event['table'] + 1}:H:C{event['column'] + 1}"
                    if event["row"] == -1
                    else f"T{event['table'] + 1}:R{event['row'] + 1}:C{event['column'] + 1}"
                ),
                event["at"],
                event["table"] + 1,
                event["row"] + 2,
                event["column"] + 1,
                event["before"],
                event["after"],
                event["note"],
                event["previous"],
                event["sha256"],
            ],
        )
    summaries = [
        [
            index,
            record.get("page"),
            record.get("role"),
            record.get("text"),
            *(record.get("bbox") or [None] * 4),
        ]
        for index, table in enumerate(result.tables, 1)
        for record in table.page_summaries
    ]
    if summaries:
        sheet = workbook.create_sheet("Page Summaries")
        strings.append(sheet, ["Table", "Page", "Role", "Source text", "X0", "Y0", "X1", "Y1"])
        for row in summaries:
            strings.append(sheet, row)
    strings.finish(workbook)
    for sheet in workbook:
        style_sheet(sheet)
    # Force text types in audit metadata as well as data cells.
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    target = Path(output)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    return target


def csv_text(value: object) -> str:
    """Return a spreadsheet-safe CSV field while preserving exact text in audit JSON."""
    text = "" if value is None else str(value)
    if re.match(r"^[\s\u0000-\u001f]*[=+\-@]", text):
        return "'" + text
    return text


def export_csv(result: Result, output: str | Path, table: int = 0) -> Path:
    """Export one selected table as UTF-8 CSV without executable spreadsheet formulas."""
    if type(table) is not int or not 0 <= table < len(result.tables):
        raise ValueError("Invalid table index.")
    target = Path(output)
    target.parent.mkdir(parents=True, exist_ok=True)
    selected = result.tables[table]
    with target.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(csv_text(value) for value in selected.headers)
        for row in selected.rows:
            writer.writerow(csv_text(value) for value in row)
    return target
