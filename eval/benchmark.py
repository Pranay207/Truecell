"""Two-stage, source-bound evaluation for genuinely held-out documents.

The extraction stage never needs access to truth.  An evaluator can provide
truth later and score the saved audit without running the extractor again.
"""

from __future__ import annotations

import hashlib
import json
import os
import platform
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from eval.score import score
from truecell.checkpoint import signature as extraction_signature
from truecell.constraints import coverage
from truecell.export import export_workbook
from truecell.models import Result
from truecell.pipeline import convert
from truecell.profiling import normalize_languages

CASE_ID = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]{0,99}$")
RUN_VERSION = 2
PROJECT_ROOT = Path(__file__).resolve().parents[1]


class ManifestError(ValueError):
    """The benchmark manifest is unsafe or incomplete."""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def source_fingerprint(root: Path = PROJECT_ROOT) -> str:
    """Hash the extraction implementation and locked dependency declarations."""
    files = [
        *sorted((root / "truecell").glob("*.py")),
        *sorted((root / "eval").glob("*.py")),
        *sorted((root / "models").glob("*.onnx")),
        root / "requirements.txt",
        root / "frontend" / "package-lock.json",
    ]
    digest = hashlib.sha256()
    for path in files:
        if not path.is_file():
            continue
        relative = path.relative_to(root).as_posix().encode("utf-8")
        digest.update(len(relative).to_bytes(4, "big"))
        digest.update(relative)
        content = path.read_bytes()
        digest.update(len(content).to_bytes(8, "big"))
        digest.update(content)
    return digest.hexdigest()


def _relative(root: Path, name: str) -> Path:
    candidate = Path(name)
    if candidate.is_absolute() or ".." in candidate.parts:
        raise ManifestError(f"Benchmark path must stay inside its pack: {name!r}")
    resolved_root = root.resolve()
    resolved = (resolved_root / candidate).resolve()
    if resolved != resolved_root and resolved_root not in resolved.parents:
        raise ManifestError(f"Benchmark path escapes its pack: {name!r}")
    return resolved


def load_manifest(path: Path) -> dict[str, Any]:
    try:
        manifest = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ManifestError(f"Cannot read benchmark manifest: {exc}") from exc
    if manifest.get("version") != 1 or not isinstance(manifest.get("cases"), list):
        raise ManifestError("Manifest must use version 1 and contain a cases list.")
    seen: set[str] = set()
    for case in manifest["cases"]:
        if not isinstance(case, dict):
            raise ManifestError("Every benchmark case must be an object.")
        identifier = case.get("id")
        if not isinstance(identifier, str) or not CASE_ID.fullmatch(identifier):
            raise ManifestError(f"Unsafe or missing case id: {identifier!r}")
        if identifier in seen:
            raise ManifestError(f"Duplicate benchmark case id: {identifier}")
        seen.add(identifier)
        if not isinstance(case.get("file"), str):
            raise ManifestError(f"Case {identifier} has no input file.")
        expected_hash = case.get("sha256")
        if not isinstance(expected_hash, str) or not re.fullmatch(r"[0-9a-f]{64}", expected_hash):
            raise ManifestError(f"Case {identifier} needs a lowercase SHA-256 hash.")
        if "truth" in case and not isinstance(case["truth"], str):
            raise ManifestError(f"Case {identifier} truth path must be a string.")
        try:
            normalize_languages(case.get("ocr_languages"))
        except ValueError as exc:
            raise ManifestError(f"Case {identifier} has invalid OCR languages: {exc}") from exc
    return manifest


def extract_manifest(manifest_path: Path, inputs: Path, output: Path) -> dict[str, Any]:
    """Convert a held-out pack while keeping truth outside the process."""
    manifest = load_manifest(manifest_path)
    output.mkdir(parents=True, exist_ok=True)
    records = []
    for case in manifest["cases"]:
        identifier = case["id"]
        source = _relative(inputs, case["file"])
        if not source.is_file():
            raise ManifestError(f"Missing benchmark input for {identifier}: {source.name}")
        observed_hash = sha256(source)
        if observed_hash != case["sha256"]:
            raise ManifestError(f"Input hash mismatch for {identifier}.")
        ocr_languages = normalize_languages(case.get("ocr_languages"))
        result = convert(
            source,
            output / f"{identifier}.xlsx",
            ocr_languages="+".join(ocr_languages),
        )
        runtime_options = {
            "verify": False,
            "force_ocr": False,
            "allow_cloud": False,
            "ocr_languages": "+".join(ocr_languages),
        }
        audit_path = output / f"{identifier}.audit.json"
        audit_path.write_text(json.dumps(result.to_dict(), indent=2), encoding="utf-8")
        workbook_path = output / f"{identifier}.xlsx"
        records.append(
            {
                "id": identifier,
                "category": case.get("category", "unspecified"),
                "language": case.get("language", "unspecified"),
                "ocr_languages": list(ocr_languages),
                "extraction_signature": extraction_signature(observed_hash, runtime_options),
                "source_sha256": observed_hash,
                "pages": len(result.pages),
                "tables": len(result.tables),
                "rows": sum(len(table.rows) for table in result.tables),
                "cells": sum(len(row) for table in result.tables for row in table.rows),
                "findings": len(result.issues),
                "failed_pages": sum(bool(page.error) for page in result.pages),
                "seconds": result.seconds,
                "evidence": coverage(result),
                "audit": audit_path.name,
                "audit_sha256": sha256(audit_path),
                "workbook": workbook_path.name,
                "workbook_sha256": sha256(workbook_path),
            }
        )
    run = {
        "run_version": RUN_VERSION,
        "suite": manifest.get("suite", manifest_path.stem),
        "manifest_sha256": sha256(manifest_path),
        "extractor_fingerprint": source_fingerprint(),
        "extractor_revision": os.getenv("TRUECELL_BUILD_REVISION"),
        "extracted_at": datetime.now(timezone.utc).isoformat(),
        "host": {
            "python": platform.python_version(),
            "system": platform.system(),
            "machine": platform.machine(),
        },
        # This records the evaluator's protocol declaration.  The separately
        # retained run.json SHA-256 seal makes later changes detectable; it is
        # not a cryptographic proof of what the extractor could access.
        "truth_available_during_extraction": False,
        "truth_protocol": "evaluator-withheld",
        "cases": records,
    }
    (output / "run.json").write_text(json.dumps(run, indent=2), encoding="utf-8")
    return run


def _verify_workbook_matches_audit(path: Path, result: Result, identifier: str) -> None:
    """Require every exported data cell to equal a fresh export of the frozen audit."""
    try:
        with tempfile.TemporaryDirectory(prefix="truecell-benchmark-") as folder:
            expected_path = Path(folder) / "expected.xlsx"
            export_workbook(result, expected_path)
            actual = expected = None
            try:
                actual = load_workbook(path, read_only=True, data_only=False)
                expected = load_workbook(expected_path, read_only=True, data_only=False)
                data_names = [f"Data {index}" for index in range(1, len(result.tables) + 1)]
                actual_data_names = [
                    name for name in actual.sheetnames if re.fullmatch(r"Data \d+", name)
                ]
                if actual_data_names != data_names:
                    raise ManifestError(
                        f"Saved workbook data sheets do not match the audit for {identifier}."
                    )
                compare_names = data_names + (
                    ["Long Text"] if "Long Text" in expected.sheetnames else []
                )
                for name in compare_names:
                    if name not in actual.sheetnames:
                        raise ManifestError(
                            f"Saved workbook is missing {name!r} required by the audit for {identifier}."
                        )
                    observed_sheet = actual[name]
                    expected_sheet = expected[name]
                    if (observed_sheet.max_row, observed_sheet.max_column) != (
                        expected_sheet.max_row,
                        expected_sheet.max_column,
                    ):
                        raise ManifestError(
                            f"Saved workbook shape does not match the audit for {identifier} ({name})."
                        )
                    for observed_row, expected_row in zip(
                        observed_sheet.iter_rows(), expected_sheet.iter_rows(), strict=True
                    ):
                        for observed, wanted in zip(observed_row, expected_row, strict=True):
                            if (
                                observed.value != wanted.value
                                or observed.data_type != wanted.data_type
                                or observed.number_format != wanted.number_format
                            ):
                                raise ManifestError(
                                    "Saved workbook values do not match the frozen audit for "
                                    f"{identifier} at {name}!{observed.coordinate}."
                                )
            finally:
                if actual is not None:
                    actual.close()
                if expected is not None:
                    expected.close()
    except ManifestError:
        raise
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as exc:
        raise ManifestError(f"Saved workbook is invalid for {identifier}: {exc}") from exc


def score_saved(
    manifest_path: Path,
    output: Path,
    truth_root: Path,
    expected_run_sha256: str,
) -> list[dict[str, Any]]:
    """Score saved audits; never rerun or mutate the extraction."""
    manifest = load_manifest(manifest_path)
    run_path = output / "run.json"
    if not run_path.is_file():
        raise ManifestError("Missing run.json; perform the extraction stage first.")
    if not re.fullmatch(r"[0-9a-f]{64}", expected_run_sha256 or ""):
        raise ManifestError("Scoring requires the evaluator's lowercase run.json SHA-256 seal.")
    if sha256(run_path) != expected_run_sha256:
        raise ManifestError("run.json changed after the evaluator recorded its SHA-256 seal.")
    try:
        run = json.loads(run_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ManifestError(f"Cannot read sealed run.json: {exc}") from exc
    if not isinstance(run, dict):
        raise ManifestError("Saved run.json root must be an object.")
    if run.get("run_version") != RUN_VERSION:
        raise ManifestError(f"Saved run must use run version {RUN_VERSION}.")
    recorded_fingerprint = run.get("extractor_fingerprint")
    if not isinstance(recorded_fingerprint, str) or not re.fullmatch(
        r"[0-9a-f]{64}", recorded_fingerprint
    ):
        raise ManifestError("Saved run is missing a valid extractor fingerprint.")
    if recorded_fingerprint != source_fingerprint():
        raise ManifestError(
            "Extraction source changed after the sealed run; use its frozen revision."
        )
    if (
        run.get("truth_available_during_extraction") is not False
        or run.get("truth_protocol") != "evaluator-withheld"
    ):
        raise ManifestError("Saved run does not declare the evaluator-withheld truth protocol.")
    if run.get("manifest_sha256") != sha256(manifest_path):
        raise ManifestError("Manifest changed after extraction; start a new benchmark run.")
    run_cases = run.get("cases", [])
    if not isinstance(run_cases, list) or any(not isinstance(record, dict) for record in run_cases):
        raise ManifestError("Saved run has an invalid cases list.")
    run_ids = [record.get("id") for record in run_cases]
    if any(
        not isinstance(identifier, str) or not CASE_ID.fullmatch(identifier)
        for identifier in run_ids
    ):
        raise ManifestError("Saved run has an invalid case id.")
    if any(
        not isinstance(record.get(name), str)
        for record in run_cases
        for name in (
            "audit",
            "workbook",
            "source_sha256",
            "audit_sha256",
            "workbook_sha256",
            "extraction_signature",
        )
    ):
        raise ManifestError("Saved run has invalid artifact metadata.")
    if any(
        not re.fullmatch(r"[0-9a-f]{64}", record[name])
        for record in run_cases
        for name in (
            "source_sha256",
            "audit_sha256",
            "workbook_sha256",
            "extraction_signature",
        )
    ):
        raise ManifestError("Saved run has invalid hash or runtime-signature metadata.")
    manifest_ids = [case["id"] for case in manifest["cases"]]
    if len(run_ids) != len(set(run_ids)) or set(run_ids) != set(manifest_ids):
        raise ManifestError("Saved run case set does not exactly match the frozen manifest.")
    records_by_id = {record["id"]: record for record in run_cases}
    metrics = []
    for case in manifest["cases"]:
        identifier = case["id"]
        record = records_by_id.get(identifier)
        if record is None or record.get("source_sha256") != case["sha256"]:
            raise ManifestError(f"Saved run is missing or mismatched for {identifier}.")
        truth_name = case.get("truth", f"{identifier}.truth.json")
        truth_path = _relative(truth_root, truth_name)
        audit_path = _relative(output, record["audit"])
        workbook_path = _relative(output, record.get("workbook", ""))
        if not truth_path.is_file() or not audit_path.is_file() or not workbook_path.is_file():
            raise ManifestError(f"Missing truth or saved extraction artifact for {identifier}.")
        if record.get("audit_sha256") != sha256(audit_path):
            raise ManifestError(f"Saved audit was changed after extraction for {identifier}.")
        if record.get("workbook_sha256") != sha256(workbook_path):
            raise ManifestError(f"Saved workbook was changed after extraction for {identifier}.")
        try:
            truth = json.loads(truth_path.read_text(encoding="utf-8"))
            audit_payload = json.loads(audit_path.read_text(encoding="utf-8"))
            if not isinstance(truth, dict) or not isinstance(audit_payload, dict):
                raise TypeError("truth and audit roots must be objects")
            result = Result.from_dict(audit_payload)
        except (OSError, json.JSONDecodeError, KeyError, TypeError, ValueError) as exc:
            raise ManifestError(f"Invalid truth or saved audit for {identifier}: {exc}") from exc
        if result.source_sha256 != case["sha256"]:
            raise ManifestError(f"Audit source hash is mismatched for {identifier}.")
        expected_languages = list(normalize_languages(case.get("ocr_languages")))
        if record.get("ocr_languages") != expected_languages:
            raise ManifestError(f"Saved OCR language routing is mismatched for {identifier}.")
        _verify_workbook_matches_audit(workbook_path, result, identifier)
        metric = score(truth, result)
        metric.update(
            {
                "id": identifier,
                "category": case.get("category", "unspecified"),
                "language": case.get("language", "unspecified"),
                "source_sha256": case["sha256"],
                "truth_sha256": sha256(truth_path),
                "audit_sha256": record["audit_sha256"],
                "workbook_sha256": record["workbook_sha256"],
                "extractor_fingerprint": recorded_fingerprint,
                "extraction_signature": record["extraction_signature"],
                "ocr_languages": expected_languages,
            }
        )
        metrics.append(metric)
    (output / "results.json").write_text(json.dumps(metrics, indent=2), encoding="utf-8")
    write_results(metrics, output / "RESULTS.md", run.get("suite", manifest_path.stem))
    return metrics


def write_results(metrics: list[dict[str, Any]], path: Path, suite: str) -> None:
    lines = [
        f"# Blind evaluation: {suite}",
        "",
        "The evaluator-withheld protocol records that truth was supplied only after extraction. Saved, source-bound audits were scored without rerunning extraction.",
        "",
        (
            f"Extractor fingerprint: `{metrics[0]['extractor_fingerprint']}`"
            if metrics
            else "Extractor fingerprint: no cases were scored."
        ),
        "",
        "| Case | Category | Language | Cells exact | Rows | Headers | Shape | Result |",
        "|---|---|---|---:|---:|---:|---|---|",
    ]
    for metric in metrics:
        lines.append(
            f"| {metric['id']} | {metric['category']} | {metric['language']} | "
            f"{metric['cell_exact']:.1%} | {metric['row_match']:.1%} | "
            f"{metric['header_match']:.1%} | {metric['shape_match']} | "
            f"{'PASS' if metric['pass'] else 'FAIL'} |"
        )
    lines += [
        "",
        f"Exact passes: **{sum(item['pass'] for item in metrics)}/{len(metrics)}**.",
        "",
        "A pass requires exact cells, headers, table shapes and all annotated injected errors. Results apply only to this frozen suite.",
    ]
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
