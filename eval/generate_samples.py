"""Generate dummy development PDFs and canonical ground truth before extraction."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib.colors import HexColor
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
WIDTH, HEIGHT = 1190, 842
HEADERS = [
    "Invoice",
    "Date",
    "Client",
    "SKU",
    "Description",
    "Qty.",
    "Rate",
    "Amt",
    "Tax",
    "Ship",
    "Region",
    "Status",
]
WIDTHS = [85, 90, 90, 75, 180, 55, 70, 80, 70, 65, 90, 100]


def draw_table(
    pdf,
    headers: list[str],
    rows: list[list[str]],
    *,
    x: float = 35,
    top: float = 750,
    widths: list[float] | None = None,
    ruled: bool = True,
    show_header: bool = True,
) -> None:
    """Draw a clean, fixed-grid table with readable dummy text."""
    widths = widths or [100] * len(headers)
    all_rows = ([headers] if show_header else []) + rows
    h = 27
    for r, row in enumerate(all_rows):
        y = top - (r + 1) * h
        if r == 0 and show_header:
            pdf.setFillColor(HexColor("#17324d"))
            pdf.rect(x, y, sum(widths), h, fill=1, stroke=0)
            pdf.setFillColor(HexColor("#ffffff"))
        else:
            pdf.setFillColor(HexColor("#182c3e"))
        offset = x
        for c, value in enumerate(row):
            pdf.setFont("Helvetica-Bold" if r == 0 and show_header else "Helvetica", 9)
            pdf.drawString(offset + 5, y + 9, str(value))
            offset += widths[c]
    if ruled:
        pdf.setStrokeColor(HexColor("#8294a5"))
        pdf.setLineWidth(0.5)
        for r in range(len(all_rows) + 1):
            pdf.line(x, top - r * h, x + sum(widths), top - r * h)
        offset = x
        for w in [0] + widths:
            offset += w
            pdf.line(offset, top, offset, top - len(all_rows) * h)


def write_truth(
    directory: Path,
    name: str,
    tables: list[dict],
    pages: int,
    injected_errors: list[dict] | None = None,
) -> None:
    """Save independently authored truth as JSON and a companion workbook."""
    truth = {
        "name": name,
        "pages": pages,
        "tables": tables,
        "injected_errors": injected_errors or [],
    }
    (directory / f"{name}.truth.json").write_text(json.dumps(truth, indent=2), encoding="utf-8")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for i, table in enumerate(tables, 1):
        sheet = workbook.create_sheet(f"Data {i}")
        sheet.append(table["headers"])
        for row in table["rows"]:
            sheet.append(row)
    if not tables:
        workbook.create_sheet("Read me").append(["No tables expected in this document."])
    workbook.save(directory / f"{name}.truth.xlsx")


def invoice_rows(count: int = 48) -> list[list[str]]:
    """Return deterministic 12-column records with identifiable row keys."""
    rows = []
    for i in range(count):
        qty, rate = i % 7 + 1, 12 + i % 11
        rows.append(
            [
                f"INV-{i + 1:04d}",
                f"2026-10-{i % 28 + 1:02d}",
                f"Client {i % 9 + 1}",
                f"00{i + 1:04d}",
                f"Widget model {i % 13 + 1}",
                str(qty),
                f"{rate:.2f}",
                f"{qty * rate:.2f}",
                f"{qty * rate * 0.1:.2f}",
                "5.00",
                "South",
                "Paid",
            ]
        )
    return rows


def invoice_sample(
    directory: Path,
    name: str = "A_invoice_register",
    *,
    pages: int = 3,
    ruled: bool = True,
    headerless: bool = False,
) -> None:
    """Create a multipage invoice register with repeated headers and page footers."""
    rows = invoice_rows(pages * 16)
    pdf = canvas.Canvas(str(directory / f"{name}.pdf"), pagesize=(WIDTH, HEIGHT), invariant=1)
    for p in range(pages):
        pdf.setFont("Helvetica-Bold", 18)
        pdf.drawString(35, 800, "Dummy invoice register")
        draw_table(
            pdf,
            HEADERS,
            rows[p * 16 : (p + 1) * 16],
            widths=WIDTHS,
            ruled=ruled,
            show_header=not headerless or p == 0,
        )
        pdf.setFont("Helvetica", 10)
        pdf.drawString(35, 35, f"Page {p + 1} of {pages}")
        pdf.showPage()
    pdf.save()
    write_truth(directory, name, [{"headers": HEADERS, "rows": rows}], pages)


def bank_sample(directory: Path) -> None:
    """Create odd-labelled transactions with a wrong total and altered balance."""
    name = "B_bank_errors"
    headers = ["Ref", "When", "Memo", "X", "Y", "Z"]
    rows = []
    balance = 1000
    for i in range(18):
        debit, credit = (11 + i, 0) if i % 3 else (0, 40 + i)
        balance += credit - debit
        displayed = balance + (17 if i == 10 else 0)
        rows.append(
            [
                f"TX-{i + 1:03d}",
                f"2026-10-{i + 1:02d}",
                f"Entry {i + 1}",
                f"{debit:.2f}",
                f"{credit:.2f}",
                f"{displayed:.2f}",
            ]
        )
    total = sum(float(row[3]) for row in rows)
    rows.append(["", "", "Total", f"{total + 9:.2f}", "", ""])
    pdf = canvas.Canvas(str(directory / f"{name}.pdf"), pagesize=(WIDTH, HEIGHT), invariant=1)
    pdf.setFont("Helvetica-Bold", 18)
    pdf.drawString(35, 800, "Dummy bank statement: deliberate errors")
    draw_table(pdf, headers, rows, widths=[120, 130, 260, 140, 140, 160])
    pdf.showPage()
    pdf.save()
    errors = [
        {"table": 1, "row": 11, "column": 6, "code": "running_balance"},
        {"table": 1, "row": 19, "column": 4, "code": "total_mismatch"},
    ]
    write_truth(directory, name, [{"headers": headers, "rows": rows}], 1, errors)


def scan_sample(
    directory: Path,
    name: str = "C_scanned_inventory",
    *,
    skew: float = 0,
    degradation: str = "",
    mixed: bool = False,
) -> None:
    """Rasterize dummy inventory pages; truth remains the original intended cells."""
    import io

    import fitz
    import numpy as np
    from PIL import Image, ImageEnhance, ImageFilter

    buffer = io.BytesIO()
    source = canvas.Canvas(buffer, pagesize=(WIDTH, HEIGHT), invariant=1)
    rows = invoice_rows(24)
    for p in range(2):
        source.setFont("Helvetica-Bold", 18)
        source.drawString(35, 800, "Dummy inventory scan")
        draw_table(source, HEADERS, rows[p * 12 : (p + 1) * 12], widths=WIDTHS)
        source.setFont("Helvetica", 10)
        source.drawString(35, 35, f"Page {p + 1} of 2")
        source.showPage()
    source.save()
    native = fitz.open(stream=buffer.getvalue(), filetype="pdf")
    result = fitz.open()
    for p in range(2):
        if mixed and p == 0:
            result.insert_pdf(native, from_page=p, to_page=p)
            continue
        pix = native[p].get_pixmap(dpi=180)
        image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        if skew:
            image = image.rotate(
                skew, expand=False, fillcolor="white", resample=Image.Resampling.BICUBIC
            )
        if degradation == "noise":
            image = ImageEnhance.Contrast(image).enhance(0.55).filter(ImageFilter.GaussianBlur(0.7))
            array = np.asarray(image).astype(float)
            rng = np.random.default_rng(17)
            image = Image.fromarray(
                np.clip(array + rng.normal(0, 5, array.shape), 0, 255).astype("uint8")
            )
        if degradation == "photo":
            import cv2

            array = np.asarray(image)
            h, w = array.shape[:2]
            matrix = cv2.getPerspectiveTransform(
                np.float32([[0, 0], [w, 0], [w, h], [0, h]]),
                np.float32([[w * 0.035, h * 0.02], [w * 0.95, 0], [w, h * 0.96], [0, h]]),
            )
            array = cv2.warpPerspective(array, matrix, (w, h), borderValue=(255, 255, 255))
            shadow = np.linspace(0.6, 1, w)[None, :, None]
            image = Image.fromarray(np.clip(array * shadow, 0, 255).astype("uint8"))
        encoded = io.BytesIO()
        image.save(encoded, format="PNG")
        page = result.new_page(width=WIDTH, height=HEIGHT)
        page.insert_image(page.rect, stream=encoded.getvalue())
    result.save(directory / f"{name}.pdf", deflate=True)
    result.close()
    native.close()
    write_truth(directory, name, [{"headers": HEADERS, "rows": rows}], 2)


def scanned_continuation_stress_sample(directory: Path) -> None:
    """Create a five-page raster table with two page-break row continuations."""
    import io

    import fitz
    from PIL import Image, ImageEnhance

    name = "F_scanned_multipage_continuations"
    headers = [
        "Batch",
        "Date",
        "Supplier",
        "Item",
        "Description",
        "UOM",
        "Qty",
        "Rate",
        "Gross",
        "Tax",
        "Zone",
        "Status",
    ]
    page_width, page_height = 900, 420
    widths = [68, 82, 78, 67, 135, 45, 45, 62, 68, 58, 72, 78]
    descriptions = [
        "Sensor module",
        "Valve kit",
        "Relay panel",
        "Cable drum",
        "Filter pack",
        "Motor guard",
        "Pump seal",
        "Gauge mount",
    ]
    rows = []
    for index in range(33):
        quantity = index % 7 + 1
        rate = 14 + index % 9
        gross = quantity * rate
        rows.append(
            [
                f"BT-{index + 1:04d}",
                f"2026-09-{index % 28 + 1:02d}",
                f"Vendor {chr(65 + index % 6)}",
                f"IT-{index + 1:04d}",
                descriptions[index % len(descriptions)],
                "ea",
                str(quantity),
                f"{rate:.2f}",
                f"{gross:.2f}",
                f"{gross * 0.1:.2f}",
                ["East", "North", "South", "West"][index % 4],
                ["Ready", "Review", "Hold"][index % 3],
            ]
        )

    def first_fragment(row: list[str]) -> list[str]:
        first_word = row[4].split(maxsplit=1)[0]
        return row[:4] + [first_word, row[5], "", "", "", "", row[10], ""]

    def second_fragment(row: list[str]) -> list[str]:
        second_word = row[4].split(maxsplit=1)[1]
        return ["", "", "", "", second_word, ""] + row[6:10] + ["", row[11]]

    # Every page contains seven physical rows. Rows 7 and 27 cross a page break.
    page_rows = [
        rows[0:6] + [first_fragment(rows[6])],
        [second_fragment(rows[6])] + rows[7:13],
        rows[13:20],
        rows[20:26] + [first_fragment(rows[26])],
        [second_fragment(rows[26])] + rows[27:33],
    ]

    native_buffer = io.BytesIO()
    source = canvas.Canvas(native_buffer, pagesize=(page_width, page_height), invariant=1)
    for page_index, physical_rows in enumerate(page_rows, 1):
        source.setFont("Helvetica-Bold", 16)
        source.drawString(20, 392, "Dummy receiving register - scanned copy")
        draw_table(source, headers, physical_rows, x=20, top=355, widths=widths)
        source.setFont("Helvetica", 10)
        source.drawString(20, 20, f"Page {page_index} of {len(page_rows)}")
        source.showPage()
    source.save()

    native = fitz.open(stream=native_buffer.getvalue(), filetype="pdf")
    scanned = fitz.open()
    for page in native:
        pix = page.get_pixmap(dpi=180, colorspace=fitz.csRGB, alpha=False)
        image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        image = ImageEnhance.Contrast(image).enhance(0.96)
        encoded = io.BytesIO()
        image.save(encoded, format="PNG", optimize=True)
        target = scanned.new_page(width=page_width, height=page_height)
        target.insert_image(target.rect, stream=encoded.getvalue())
    scanned.save(directory / f"{name}.pdf", deflate=True)
    scanned.close()
    native.close()
    write_truth(directory, name, [{"headers": headers, "rows": rows}], len(page_rows))


def layout_samples(directory: Path) -> None:
    """Create distinct hard-layout cases with explicit intended output shapes."""
    # Two unrelated ruled tables on the same page.
    name = "D_two_tables"
    pdf = canvas.Canvas(str(directory / f"{name}.pdf"), pagesize=(WIDTH, HEIGHT), invariant=1)
    h1, h2 = ["Code", "Count", "Price"], ["Branch", "In", "Out", "Net"]
    r1 = [[f"P{i}", str(i + 1), str(10 + i)] for i in range(5)]
    r2 = [[f"B{i}", str(40 + i), str(10 + i), "30"] for i in range(5)]
    draw_table(pdf, h1, r1, top=750, widths=[200] * 3)
    draw_table(pdf, h2, r2, top=450, widths=[150] * 4)
    pdf.showPage()
    pdf.save()
    write_truth(directory, name, [{"headers": h1, "rows": r1}, {"headers": h2, "rows": r2}], 1)
    # Merged two-level header; body text wraps within ruled cells.
    name = "D_merged_wrapped"
    pdf = canvas.Canvas(str(directory / f"{name}.pdf"), pagesize=(WIDTH, HEIGHT), invariant=1)
    x, top, cw, rh = 35, 750, 180, 42
    headers = [
        "Product | Code",
        "Product | Description",
        "Movement | Units",
        "Movement | Value",
    ]
    rows = [
        ["A01", "Long product name", "12", "36.00"],
        ["A02", "Wrapped description", "15", "45.00"],
        ["A03", "Single", "10", "30.00"],
    ]
    pdf.setFont("Helvetica", 11)
    pdf.rect(x, top - 2 * rh, 4 * cw, 2 * rh)
    pdf.line(x, top - rh, x + 4 * cw, top - rh)
    pdf.line(x + 2 * cw, top, x + 2 * cw, top - 2 * rh)
    for c in (1, 3):
        pdf.line(x + c * cw, top - rh, x + c * cw, top - 2 * rh)
    pdf.drawString(x + 5, top - 26, "Product")
    pdf.drawString(x + 2 * cw + 5, top - 26, "Movement")
    for c, label in enumerate(["Code", "Description", "Units", "Value"]):
        pdf.drawString(x + c * cw + 5, top - rh - 26, label)
    for r, row in enumerate(rows):
        y = top - (r + 3) * rh
        for c, value in enumerate(row):
            pdf.rect(x + c * cw, y, cw, rh)
            if c == 1 and r < 2:
                split = value.rsplit(" ", 1)
                pdf.drawString(x + c * cw + 5, y + 25, split[0])
                pdf.drawString(x + c * cw + 5, y + 10, split[1])
            else:
                pdf.drawString(x + c * cw + 5, y + 17, value)
    pdf.showPage()
    pdf.save()
    write_truth(directory, name, [{"headers": headers, "rows": rows}], 1)
    # A physical row split at a page break: both fragments must be retained if uncertain.
    name = "D_split_row"
    pdf = canvas.Canvas(str(directory / f"{name}.pdf"), pagesize=(WIDTH, HEIGHT), invariant=1)
    headers = ["Ref", "Description", "Units", "Value"]
    draw_table(
        pdf,
        headers,
        [["A1", "Complete", "2", "10"], ["A2", "Long item", "", ""]],
        widths=[150, 350, 150, 150],
    )
    pdf.showPage()
    draw_table(
        pdf,
        headers,
        [["", "continued", "3", "15"], ["A3", "Next", "4", "20"]],
        widths=[150, 350, 150, 150],
    )
    pdf.showPage()
    pdf.save()
    write_truth(
        directory,
        name,
        [
            {
                "headers": headers,
                "rows": [
                    ["A1", "Complete", "2", "10"],
                    ["A2", "Long item continued", "3", "15"],
                    ["A3", "Next", "4", "20"],
                ],
            }
        ],
        2,
    )


def main() -> None:
    """Write the baseline dummy development set."""
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=ROOT / "samples")
    parser.add_argument("--hard", action="store_true")
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    invoice_sample(args.output)
    bank_sample(args.output)
    scan_sample(args.output)
    if args.hard:
        invoice_sample(args.output, "D_borderless", ruled=False)
        invoice_sample(args.output, "D_headerless", headerless=True)
        invoice_sample(args.output, "D_stress_50pages", pages=50)
        scan_sample(args.output, "D_skew_5deg", skew=5)
        scan_sample(args.output, "D_blur_noise", degradation="noise")
        scan_sample(args.output, "D_phone_photo", degradation="photo")
        scan_sample(args.output, "D_mixed", mixed=True)
        scanned_continuation_stress_sample(args.output)
        layout_samples(args.output)
    print(f"Generated dummy development documents in {args.output}")


if __name__ == "__main__":
    main()
