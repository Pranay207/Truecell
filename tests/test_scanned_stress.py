"""Exact evidence for the original multi-page scanned continuation fixture."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import fitz
from openpyxl import load_workbook

from eval.score import score
from truecell.models import Result, Table
from truecell.stitch import stitch

ROOT = Path(__file__).resolve().parents[1]
NAME = "F_scanned_multipage_continuations"


def load_truth() -> dict:
    return json.loads((ROOT / "samples" / f"{NAME}.truth.json").read_text(encoding="utf-8"))


def first_fragment(row: list[str]) -> list[str]:
    return row[:4] + [row[4].split(maxsplit=1)[0], row[5], "", "", "", "", row[10], ""]


def second_fragment(row: list[str]) -> list[str]:
    return ["", "", "", "", row[4].split(maxsplit=1)[1], ""] + row[6:10] + ["", row[11]]


def test_scanned_stress_fixture_is_raster_only_wide_and_truth_workbook_is_exact():
    truth = load_truth()
    table = truth["tables"][0]
    assert truth["pages"] == 5
    assert len(table["headers"]) == 12
    assert len(table["rows"]) == 33

    with fitz.open(ROOT / "samples" / f"{NAME}.pdf") as document:
        assert len(document) == 5
        assert all(not page.get_text().strip() for page in document)
        assert all(len(page.get_images(full=True)) == 1 for page in document)

    sheet = load_workbook(ROOT / "samples" / f"{NAME}.truth.xlsx", data_only=True)["Data 1"]
    values = [list(row) for row in sheet.iter_rows(values_only=True)]
    assert values == [table["headers"], *table["rows"]]


def test_scanned_stress_join_reconstructs_both_split_rows_exactly():
    truth = load_truth()
    headers = truth["tables"][0]["headers"]
    rows = truth["tables"][0]["rows"]
    physical_pages = [
        rows[0:6] + [first_fragment(rows[6])],
        [second_fragment(rows[6])] + rows[7:13],
        rows[13:20],
        rows[20:26] + [first_fragment(rows[26])],
        [second_fragment(rows[26])] + rows[27:33],
    ]
    fragments = [
        Table(headers, page_rows, page, (20, 65, 878, 355), "ocr-regression", 0.99)
        for page, page_rows in enumerate(physical_pages, 1)
    ]

    tables, issues = stitch(fragments)
    metric = score(truth, Result(f"{NAME}.pdf", tables, issues))
    assert metric["pass"]
    assert len(tables) == 1
    assert len(tables[0].rows) == 33
    assert [issue.code for issue in issues] == ["split_row_join", "split_row_join"]


def test_verified_scanned_stress_audit_matches_every_truth_cell():
    source = ROOT / "samples" / f"{NAME}.pdf"
    audit = json.loads((ROOT / "outputs" / f"{NAME}.audit.json").read_text(encoding="utf-8"))
    result = Result.from_dict(audit)
    metric = score(load_truth(), result)

    assert result.source_sha256 == hashlib.sha256(source.read_bytes()).hexdigest()
    assert metric["pass"]
    assert metric["cell_exact"] == 1
    assert len(result.tables) == 1
    assert (len(result.tables[0].rows), len(result.tables[0].headers)) == (33, 12)
    assert all(
        page.route == "ocr" and page.engine in {"tesseract", "rapidocr"} for page in result.pages
    )
    assert all(
        {attempt.engine for attempt in page.attempts} >= {"tesseract", "rapidocr"}
        for page in result.pages
    )
    assert sum(issue.code == "split_row_join" for issue in result.issues) == 2
