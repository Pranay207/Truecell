"""Excel export must preserve awkward source text and reject misaddressed edits."""

from copy import deepcopy
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from truecell.export import EXCEL_TEXT_LIMIT, csv_text, excel_text, export_csv, export_workbook
from truecell.models import Issue, Result, Table
from truecell.review import apply_edit, complete_review, review_queue, set_item_review


def check_every_review_item(result):
    for item in review_queue(result, limit=None)["items"]:
        if item["status"] == "pending":
            result = set_item_review(result, item["id"], True)
    return result


def decode_text(encoded):
    """Decode the documented left-to-right XLSX text convention."""
    decoded = []
    position = 0
    while position < len(encoded):
        if encoded[position : position + 2] == "\\\\":
            decoded.append("\\")
            position += 2
        elif encoded[position : position + 2] == "\\u":
            decoded.append(chr(int(encoded[position + 2 : position + 6], 16)))
            position += 6
        else:
            decoded.append(encoded[position])
            position += 1
    return "".join(decoded)


@pytest.mark.parametrize(
    "value",
    [
        "actual\x00control",
        r"literal\u0000sequence",
        r"C:\Users\dummy\report.pdf",
        "before\x01\x0b\x0c\x1fafter",
        "tab\tline\nreturn\r",
        "literal_x0000_escape_x005F_",
        "surrogate\ud800noncharacter\ufffe",
    ],
)
def test_excel_encoding_is_reversible_without_escape_collisions(value):
    assert decode_text(excel_text(value)) == value


def test_controls_in_data_header_evidence_and_review_metadata_export(tmp_path):
    header = "Heading\x00" + r"\u0000"
    original = "Original\x01" + r"\u0001"
    corrected = "Reviewed\x02" + r"\u0002"
    reason = "Checked\x03source"
    review_note = "Accepted\x04findings"
    table = Table([header], [[original]], 1, (0, 0, 10, 10), "native")
    table.header_sources = [[{"page": 1, "bbox": [0, 0, 10, 2], "text": header}]]
    table.cell_sources = [[[{"page": 1, "bbox": [0, 2, 10, 10], "text": original}]]]
    result = apply_edit(Result("dummy\x05.pdf", [table]), 0, 0, 0, corrected, reason)
    result = complete_review(check_every_review_item(result), True, review_note)
    before = deepcopy(result.to_dict())

    workbook = load_workbook(export_workbook(result, tmp_path / "controls.xlsx"))

    assert decode_text(workbook["Data 1"]["A1"].value) == header
    assert decode_text(workbook["Data 1"]["A2"].value) == corrected
    assert decode_text(workbook["Cell Evidence"]["K2"].value) == header
    assert decode_text(workbook["Cell Evidence"]["K3"].value) == original
    assert decode_text(workbook["Correction History"]["E2"].value) == original
    assert decode_text(workbook["Correction History"]["F2"].value) == corrected
    assert decode_text(workbook["Correction History"]["G2"].value) == reason
    verification = {row[0].value: row[1].value for row in workbook["Verification"]}
    assert decode_text(verification["Review note"]) == review_note
    assert decode_text(workbook["Read me"]["B2"].value) == result.source
    assert result.to_dict() == before
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)


def test_reviewed_export_is_gated_and_labels_draft_and_reviewed(tmp_path):
    table = Table(["ID"], [["001"]], 1, (0, 0, 10, 10), "native")
    table.header_sources = [[{"page": 1, "bbox": [0, 0, 10, 2], "text": "ID"}]]
    table.cell_sources = [[[{"page": 1, "bbox": [0, 2, 10, 10], "text": "001"}]]]
    result = Result("dummy.pdf", [table])
    draft = load_workbook(export_workbook(result, tmp_path / "draft.xlsx"))
    values = {row[0].value: row[1].value for row in draft["Verification"]}
    assert values["Export type"] == "Draft Excel"
    with pytest.raises(ValueError, match="Reviewed Excel"):
        export_workbook(result, tmp_path / "blocked.xlsx", reviewed=True)
    reviewed = complete_review(result, True, "Checked the only source cell.")
    final = load_workbook(export_workbook(reviewed, tmp_path / "reviewed.xlsx", reviewed=True))
    values = {row[0].value: row[1].value for row in final["Verification"]}
    assert values["Export type"] == "Reviewed Excel"


def test_csv_export_is_utf8_and_formula_safe(tmp_path):
    result = Result(
        "dummy.pdf",
        [
            Table(
                ["Name", "Formula"],
                [["Jos\u00e9", "=2+2"], ["x", "+cmd"]],
                1,
                (0, 0, 1, 1),
                "native",
            )
        ],
    )
    path = export_csv(result, tmp_path / "table.csv")
    text = path.read_text(encoding="utf-8-sig")
    assert "Jos\u00e9" in text and "'=2+2" in text and "'+cmd" in text
    assert csv_text("ordinary") == "ordinary"


def test_excel_limit_keeps_complete_values_in_numbered_chunks(tmp_path):
    long_text = "source\x00" + (r"literal\u0000" * 4000)
    table = Table([long_text], [[long_text]], 1, (0, 0, 10, 10), "native")
    result = Result(long_text, [table])
    before = deepcopy(result.to_dict())
    workbook = load_workbook(export_workbook(result, tmp_path / "long.xlsx"))

    assert "this cell is a preview" in workbook["Data 1"]["A1"].value
    assert "this cell is a preview" in workbook["Data 1"]["A2"].value
    preserved = {}
    for _, source_sheet, source_cell, chunk, chunks, encoded in workbook["Long Text"].iter_rows(
        min_row=2, values_only=True
    ):
        preserved.setdefault((source_sheet, source_cell), []).append((chunk, chunks, encoded))
    for coordinate in [("Data 1", "A1"), ("Data 1", "A2"), ("Read me", "B2")]:
        parts = sorted(preserved[coordinate])
        assert len(parts) == parts[0][1]
        assert decode_text("".join(part[2] for part in parts)) == long_text
    assert all(
        len(cell.value) <= EXCEL_TEXT_LIMIT
        for sheet in workbook
        for row in sheet
        for cell in row
        if isinstance(cell.value, str)
    )
    assert result.to_dict() == before


def test_tiny_numbers_do_not_underflow_or_display_as_rounded_zero(tmp_path):
    subnormal = "0." + "0" * 309 + "1"
    underflow = "0." + "0" * 399 + "1"
    values = [
        "0.0000000000000001",
        "-0.0000000000000001",
        subnormal,
        underflow,
        "0.00000000000001%",
        "18.125%",
        "0",
    ]
    table = Table(["Value"], [[value] for value in values], 1, (0, 0, 10, 10), "native")
    workbook = load_workbook(export_workbook(Result("dummy", [table]), tmp_path / "tiny.xlsx"))
    cells = [workbook["Data 1"].cell(row, 1) for row in range(2, len(values) + 2)]
    assert cells[0].value == 1e-16 and cells[0].number_format == "General"
    assert cells[1].value == -1e-16 and cells[1].number_format == "General"
    for index in [2, 3, 4]:
        assert cells[index].value == values[index] and cells[index].data_type == "s"
        assert Decimal(values[index].rstrip("%")) != 0
    assert cells[5].value == 0.18125 and cells[5].number_format == "0.000%"
    assert cells[6].value == 0 and cells[6].data_type == "n"


def test_escaped_literals_keep_identifier_text_and_raw_values(tmp_path):
    values = ["000123", "123456789012345678", r"ID\u0000", r"ID\part"]
    table = Table(["Identifier"], [[value] for value in values], 1, (0, 0, 10, 10), "native")
    result = Result("dummy", [table])
    before = deepcopy(result.to_dict())
    workbook = load_workbook(export_workbook(result, tmp_path / "identifiers.xlsx"))
    cells = [workbook["Data 1"].cell(row, 1) for row in range(2, len(values) + 2)]
    assert all(cell.data_type == "s" for cell in cells)
    assert cells[0].number_format == cells[1].number_format == "@"
    assert [decode_text(cell.value) for cell in cells] == values
    assert result.to_dict() == before


def test_shared_body_styles_preserve_numeric_formats_and_findings(tmp_path):
    table = Table(
        ["Identifier", "Amount", "Rate"],
        [["000123", "12.50", "18.125%"]],
        1,
        (0, 0, 10, 10),
        "native",
    )
    result = Result("dummy", [table], [Issue("warning", 1, 1, 1, "Check identifier", table=1)])
    workbook = load_workbook(export_workbook(result, tmp_path / "styles.xlsx"))
    sheet = workbook["Data 1"]
    assert sheet["A1"].font.bold
    assert sheet["A2"].fill.fgColor.rgb.endswith("FFF0B3")
    assert sheet["A2"].number_format == "@"
    assert sheet["B2"].number_format == "#,##0.00"
    assert sheet["C2"].number_format == "0.000%"
    assert sheet["A2"].font.name == "Calibri" and sheet["A2"].font.sz == 11
    assert sheet["A2"].alignment.wrap_text and sheet["A2"].alignment.vertical == "top"


@pytest.mark.parametrize(
    "table,row,column",
    [
        (-1, 0, 0),
        (1, 0, 0),
        (0, -2, 0),
        (0, 1, 0),
        (0, 0, -1),
        (0, 0, 1),
        (0, -1, 1),
        (False, 0, 0),
        (0, False, 0),
        (0, 0, False),
    ],
)
def test_invalid_edit_addresses_raise_before_changing_data(table, row, column):
    result = Result("dummy", [Table(["A"], [["original"]], 1, (0, 0, 10, 10), "native")])
    before = deepcopy(result.to_dict())
    with pytest.raises(ValueError, match="Invalid"):
        apply_edit(result, table, row, column, "changed", "Checked original")
    assert result.to_dict() == before
