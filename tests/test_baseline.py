"""Meaningful initial regression checks for routing, preservation and typed export."""

import pytest
from openpyxl import load_workbook

from eval.generate_samples import invoice_sample
from eval.score import score
from truecell.export import export_workbook
from truecell.models import Result, Table
from truecell.pipeline import convert
from truecell.router import InputError, open_document
from truecell.stitch import stitch
from truecell.table import normalize
from truecell.values import is_identifier, number


def test_digital_end_to_end(tmp_path):
    import json

    invoice_sample(tmp_path)
    result = convert(tmp_path / "A_invoice_register.pdf", tmp_path / "result.xlsx")
    truth = json.loads((tmp_path / "A_invoice_register.truth.json").read_text())
    assert score(truth, result)["pass"]
    assert len(result.tables[0].provenance) == 48


def test_export_formula_injection_and_identifiers(tmp_path):
    table = Table(
        ["ID", "Value"],
        [["00123", "=2+2"], ["12345678901234567", "13.50"]],
        1,
        (0, 0, 10, 10),
        "native",
    )
    output = tmp_path / "safe.xlsx"
    export_workbook(Result("=unsafe", [table]), output)
    wb = load_workbook(output)
    assert wb["Data 1"]["A2"].value == "00123"
    assert wb["Data 1"]["A3"].value == "12345678901234567"
    assert wb["Data 1"]["B2"].data_type == "s"
    assert wb["Read me"]["B2"].data_type == "s"
    assert wb["Data 1"]["B3"].value == 13.5


def test_duplicate_business_rows_never_dropped():
    a = Table(["ID", "Amt"], [["A", "12"], ["A", "12"]], 1, (0, 0, 100, 100), "native")
    b = Table(["ID", "Amt"], [["A", "12"]], 2, (0, 0, 100, 100), "native")
    joined, _ = stitch([a, b])
    assert len(joined) == 1 and len(joined[0].rows) == 3
    assert joined[0].provenance[-1] == [(2, 1)]


def test_headerless_first_row_preserved():
    table = normalize([["A1", "20"], ["A2", "30"]], 1, (0, 0, 10, 10), "native")
    assert table.rows[0] == ["A1", "20"]
    assert not table.header_detected


@pytest.mark.parametrize(
    "value,expected",
    [
        ("1,234.50", "1234.50"),
        ("(12.3)", "-12.3"),
        ("₹25.10", "25.10"),
        ("12,34", None),
        ("1O.00", None),
        ("", None),
    ],
)
def test_decimal_parser(value, expected):
    actual = number(value)
    assert (str(actual) if actual is not None else None) == expected


@pytest.mark.parametrize(
    "value,expected",
    [("00123", True), ("1234567890123456", True), ("12", False), ("0.2", False)],
)
def test_identifiers(value, expected):
    assert is_identifier(value) == expected


def test_input_missing(tmp_path):
    with pytest.raises(InputError, match="not found"):
        with open_document(tmp_path / "missing.pdf"):
            pass
