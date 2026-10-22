"""Unicode-native semantics and honest local OCR capability evidence."""

import json
from pathlib import Path

from openpyxl import load_workbook

from truecell.export import export_workbook
from truecell.models import Result, Table
from truecell.ocr import engine_language_coverage, order_engines
from truecell.profiling import detect_scripts, normalize_languages
from truecell.semantics import infer_column_semantics

FIXTURE = Path(__file__).parent / "fixtures" / "multilingual_tables.json"


def _source(text: str, page: int, column: int, row: int) -> dict:
    return {
        "page": page,
        "bbox": [column * 90, row * 20, (column + 1) * 90, (row + 1) * 20],
        "text": text,
        "engine": "native-lines",
        "confidence": None,
        "coordinates": "page points",
    }


def test_genuine_hindi_telugu_fixture_roundtrips_semantics_and_excel(tmp_path):
    fixtures = json.loads(FIXTURE.read_text(encoding="utf-8"))
    tables = []
    for page, fixture in enumerate(fixtures, 1):
        table = Table(
            fixture["headers"],
            fixture["rows"],
            page,
            (0, 0, 450, 100),
            "native-lines",
        )
        table.header_sources = [
            [_source(value, page, column, 0)] for column, value in enumerate(table.headers)
        ]
        table.cell_sources = [
            [[_source(value, page, column, row + 1)] for column, value in enumerate(values)]
            for row, values in enumerate(table.rows)
        ]
        assert [item.role for item in infer_column_semantics(table)] == fixture["roles"]
        text = " ".join(table.headers + [value for row in table.rows for value in row])
        assert fixture["script"] in detect_scripts(text)
        tables.append(table)

    result = Result("multilingual-native.json", tables=tables)
    restored = Result.from_dict(json.loads(json.dumps(result.to_dict(), ensure_ascii=False)))
    assert restored.tables[0].rows[0][2] == "अनिता शर्मा"
    assert restored.tables[1].rows[0][2] == "అనిత"

    workbook = load_workbook(export_workbook(restored, tmp_path / "multilingual.xlsx"))
    assert workbook["Data 1"]["C2"].value == "अनिता शर्मा"
    assert workbook["Data 2"]["C2"].value == "అనిత"
    evidence = workbook["Cell Evidence"]
    headers = {cell.value: cell.column for cell in evidence[1]}
    assert any(
        evidence.cell(row, headers["Original source reading"]).value == "అనిత"
        and evidence.cell(row, headers["Confidence type"]).value == "native_text"
        for row in range(2, evidence.max_row + 1)
    )


def test_language_aliases_and_engine_coverage_report_missing_packs_honestly(monkeypatch):
    import truecell.ocr as ocr

    requested = normalize_languages("English, Hindi, Telugu")
    assert requested == ("eng", "hin", "tel")
    monkeypatch.setattr(ocr, "tesseract_languages", lambda: ("eng", "hin"))
    assert engine_language_coverage("rapidocr", requested) == ("eng",)
    assert engine_language_coverage("tesseract", requested) == ("eng", "hin")
    assert order_engines(["rapidocr", "tesseract"], requested) == ["tesseract", "rapidocr"]
    supported = {
        language
        for engine in ("rapidocr", "tesseract")
        for language in engine_language_coverage(engine, requested)
    }
    assert [language for language in requested if language not in supported] == ["tel"]
