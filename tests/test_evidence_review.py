"""End-user evidence, diagnosis, persistence and review safety scenarios."""

import json

import pytest
from openpyxl import load_workbook

from eval.generate_samples import bank_sample, invoice_sample
from truecell.checkpoint import Checkpoints
from truecell.constraints import analyze, coverage
from truecell.models import Result
from truecell.pipeline import convert
from truecell.review import (
    apply_edit,
    cell_evidence,
    complete_review,
    confidence_evidence,
    evidence_coverage,
    evidence_status,
    fingerprint,
    history_valid,
    quality_state,
    review_current,
    review_queue,
    set_item_review,
)
from truecell.workspace import create_job, job_path, load_result, save_result


@pytest.fixture
def bank(tmp_path):
    bank_sample(tmp_path)
    return convert(tmp_path / "B_bank_errors.pdf")


def check_every_review_item(result):
    """Model the required per-item source check before the final attestation."""
    current = result
    for item in review_queue(current, limit=None)["items"]:
        if item["status"] == "pending":
            current = set_item_review(current, item["id"], True)
    return current


def test_diagnosis_uses_both_balance_transitions(bank):
    proposals = [p for p in bank.proposals if p["support_count"] >= 2]
    assert len(proposals) == 1
    assert (proposals[0]["row"], proposals[0]["column"], proposals[0]["proposed"]) == (
        10,
        5,
        "1064.00",
    )
    assert bank.tables[0].rows[10][5] == "1081.00"


def test_diagnosis_does_not_depend_on_header_labels(bank):
    bank.tables[0].headers = [f"Unknown{i}" for i in range(6)]
    analyze(bank)
    assert any(p["proposed"] == "1064.00" and p["support_count"] == 2 for p in bank.proposals)


def test_every_native_cell_has_a_real_source_box(bank):
    assert coverage(bank)["cells_with_source_boxes"] == coverage(bank)["cells"]
    source = bank.tables[0].cell_sources[10][5][0]
    assert source["text"] == "1081.00" and source["page"] == 1
    assert source["bbox"][2] > source["bbox"][0]
    assert source["confidence"] is None
    report = evidence_coverage(bank)
    assert report["nonempty_cells"] and report["unlocated"] == 0 and report["complete"]


def test_cell_evidence_envelope_distinguishes_native_scores_and_tracks_history(bank):
    detail = cell_evidence(bank, 0, 10, 5)
    assert detail["evidence_id"] == "T1:R11:C6"
    assert detail["location"]["evidence_state"] == "exact_source_box"
    assert detail["confidence"] == {
        "kind": "native_text",
        "score": None,
        "label": "Native PDF text (no OCR score)",
        "engine": "native-lines",
        "meaning": "Text and geometry came directly from the PDF text layer.",
    }
    assert detail["original_reading"] == "1081.00"
    assert detail["sources"][0]["confidence_evidence"]["kind"] == "native_text"

    edited = apply_edit(
        bank,
        0,
        10,
        5,
        "1064.00",
        "Confirmed the source reading and corrected the arithmetic value.",
    )
    changed = cell_evidence(edited, 0, 10, 5)
    assert changed["value"] == "1064.00"
    assert changed["original_value"] == "1081.00"
    assert changed["original_reading"] == "1081.00"
    assert changed["correction_count"] == 1
    assert changed["correction_history"][0]["before"] == "1081.00"


def test_multilevel_header_evidence_reports_the_complete_source_path(bank):
    bank.tables[0].headers[0] = "Identity | Code"
    bank.tables[0].header_sources[0] = [
        {
            "page": 1,
            "bbox": [0, 0, 100, 20],
            "text": "Identity",
            "engine": "native-lines",
            "role": "merged_header_path",
        },
        {
            "page": 1,
            "bbox": [0, 20, 50, 40],
            "text": "Code",
            "engine": "native-lines",
            "role": "merged_header_path",
        },
    ]

    detail = cell_evidence(bank, 0, -1, 0)

    assert detail["original_reading"] == "Identity | Code"
    assert detail["location"]["evidence_state"] == "exact_source_box"


def test_cell_evidence_retains_competing_reading_and_explicit_page_only_state(bank):
    source = bank.tables[0].cell_sources[0][0][0]
    bank.tables[0].cell_sources[0][0].append(
        {
            "page": source["page"],
            "bbox": source["bbox"],
            "text": "TXN-000X",
            "engine": "ocr-tesseract",
            "confidence": 0.71,
            "role": "comparison",
        }
    )
    detail = cell_evidence(bank, 0, 0, 0)
    assert detail["alternative_readings"][0]["text"] == "TXN-000X"
    assert detail["alternative_readings"][0]["confidence"]["kind"] == "ocr_score"

    bank.tables[0].cell_sources[0][0] = []
    page_only = cell_evidence(bank, 0, 0, 0)
    assert page_only["location"]["evidence_state"] == "page_only_review"
    assert page_only["location"]["page"] == 1
    assert page_only["confidence"]["kind"] == "not_reported"


def test_comparison_box_does_not_locate_the_current_working_value(bank):
    primary = dict(bank.tables[0].cell_sources[0][0][0])
    primary["bbox"] = None
    comparison = {
        **bank.tables[0].cell_sources[0][0][0],
        "text": "TXN-ALTERNATIVE",
        "role": "comparison",
    }
    bank.tables[0].cell_sources[0][0] = [primary, comparison]

    detail = cell_evidence(bank, 0, 0, 0)
    assert detail["location"]["evidence_state"] == "page_only_review"
    assert detail["location"]["located"] is False
    assert detail["alternative_readings"][0]["text"] == "TXN-ALTERNATIVE"
    assert evidence_coverage(bank)["unlocated"] == 1
    assert any(item["id"] == "evidence-0-0-0" for item in review_queue(bank, limit=None)["items"])


def test_confidence_evidence_clamps_engine_scores_without_calling_them_accuracy():
    detail = confidence_evidence({"engine": "ocr-test", "confidence": 1.2})
    assert detail["score"] == 1.0 and detail["kind"] == "ocr_score"
    assert "not measured accuracy" in detail["meaning"]


def test_correction_preserves_source_and_invalidates_review(bank):
    reviewed = complete_review(
        check_every_review_item(bank),
        True,
        "Checked the source and acknowledged its wrong values.",
    )
    assert review_current(reviewed)
    edited = apply_edit(
        reviewed,
        0,
        10,
        5,
        "1064.00",
        "Derived arithmetic correction; source still reads 1081.00.",
    )
    assert not review_current(edited)
    assert edited.tables[0].cell_sources[10][5][0]["text"] == "1081.00"
    assert bank.tables[0].rows[10][5] == "1081.00"
    assert not any(i.code == "running_balance" for i in edited.issues)
    assert history_valid(edited)


def test_tampered_data_invalidates_review(bank):
    reviewed = complete_review(check_every_review_item(bank), True, "Checked all source pages.")
    reviewed.tables[0].rows[0][5] = "999"
    assert not review_current(reviewed)


def test_header_edit_preserves_original_reading_and_resets_review(bank):
    reviewed = complete_review(check_every_review_item(bank), True, "Checked the source.")
    original = bank.tables[0].header_sources[0][0]["text"]
    edited = apply_edit(
        reviewed, 0, -1, 0, "Transaction reference", "Clearer label after source review."
    )
    assert edited.tables[0].headers[0] == "Transaction reference"
    assert edited.tables[0].header_sources[0][0]["text"] == original
    assert edited.history[-1]["row"] == -1 and not review_current(edited)


def test_tampered_history_is_rejected(bank):
    edited = apply_edit(bank, 0, 0, 2, "New note", "Checked source text.")
    edited.history[0]["after"] = "Different note"
    assert not history_valid(edited)
    with pytest.raises(ValueError, match="history"):
        apply_edit(edited, 0, 1, 2, "x", "Source checked.")


@pytest.mark.parametrize("checked,note", [(False, "Checked"), (True, "")])
def test_review_requires_explicit_source_attestation(bank, checked, note):
    with pytest.raises(ValueError):
        complete_review(bank, checked, note)


def test_audit_roundtrip_retains_boxes_scores_and_constraints(bank):
    bank.tables[0].cell_confidence[(0, 0)] = 0.72
    restored = Result.from_dict(json.loads(json.dumps(bank.to_dict())))
    assert restored.tables[0].cell_confidence[(0, 0)] == 0.72
    assert restored.tables[0].cell_sources == bank.tables[0].cell_sources
    assert restored.constraints == bank.constraints


def test_priority_review_queue_persists_cell_checks_without_changing_snapshot(bank):
    bank.tables[0].cell_confidence[(0, 0)] = 0.72
    queue = review_queue(bank)
    uncertain = next(item for item in queue["items"] if item["kind"] == "confidence")
    assert uncertain["table"] == 0 and uncertain["row"] == 0 and uncertain["column"] == 0
    assert uncertain["status"] == "pending" and queue["pending"] == queue["visible"]

    before = fingerprint(bank)
    checked = set_item_review(bank, uncertain["id"], True)
    updated = review_queue(checked)
    assert fingerprint(checked) == before
    assert (
        next(item for item in updated["items"] if item["id"] == uncertain["id"])["status"]
        == "checked"
    )
    assert updated["checked"] == 1

    reopened = set_item_review(checked, uncertain["id"], False)
    assert review_queue(reopened)["checked"] == 0
    with pytest.raises(ValueError, match="current extraction"):
        set_item_review(bank, "not-a-real-review-item", True)


def test_source_attestation_requires_each_priority_item(bank):
    bank.tables[0].cell_confidence[(0, 0)] = 0.72
    with pytest.raises(ValueError, match="individually"):
        complete_review(bank, True, "Premature attestation.")
    reviewed = complete_review(
        check_every_review_item(bank),
        True,
        "Checked every priority item against its source.",
    )
    queue = review_queue(reviewed)
    assert queue["visible"] > 0 and queue["pending"] == 0
    assert queue["checked"] == queue["visible"] and review_current(reviewed)
    reopened = set_item_review(reviewed, queue["items"][0]["id"], False)
    assert not review_current(reopened)
    assert review_queue(reopened)["pending"] == 1


def test_review_queue_is_paginated_without_hiding_remaining_items(bank):
    bank.tables[0].cell_confidence = {
        (row, column): 0.5
        for row, values in enumerate(bank.tables[0].rows)
        for column, _ in enumerate(values)
    }
    first = review_queue(bank, limit=10)
    assert first["visible"] == 10 and first["has_more"]
    assert first["remaining"] == first["filtered_total"] - 10
    second = review_queue(bank, offset=first["next_offset"], limit=10)
    assert not ({item["id"] for item in first["items"]} & {item["id"] for item in second["items"]})


def test_nonempty_cell_without_box_has_explicit_evidence_reason(bank):
    bank.tables[0].cell_sources[0][0] = []
    location = evidence_status(bank.tables[0].rows[0][0], [], 1)
    assert location["status"] == "unlocated" and "No cell-level" in location["reason"]
    report = evidence_coverage(bank)
    assert report["unlocated"] == 1 and not report["complete"]
    queue = review_queue(bank, limit=None)
    assert any(item["id"] == "evidence-0-0-0" for item in queue["items"])


def test_quality_gate_separates_rescan_review_and_reviewed_states(bank):
    assert quality_state(Result("empty.pdf"))["status"] == "needs_rescan"
    assert quality_state(bank)["status"] == "needs_review"
    reviewed = complete_review(check_every_review_item(bank), True, "Checked every source item.")
    assert quality_state(reviewed)["status"] == "reviewed"


def test_workspace_baseline_is_immutable(bank, tmp_path):
    identifier, _ = create_job(
        "../../dummy.pdf", (tmp_path / "B_bank_errors.pdf").read_bytes(), tmp_path
    )
    save_result(identifier, bank, tmp_path)
    edited = apply_edit(bank, 0, 10, 5, "1064.00", "Arithmetic correction after source review.")
    save_result(identifier, edited, tmp_path)
    assert load_result(identifier, tmp_path).tables[0].rows[10][5] == "1064.00"
    assert load_result(identifier, tmp_path, baseline=True).tables[0].rows[10][5] == "1081.00"
    assert len(list(tmp_path.rglob("source.pdf"))) == 1


def test_corrupt_saved_payload_rejected(bank, tmp_path):
    identifier, _ = create_job("dummy.pdf", (tmp_path / "B_bank_errors.pdf").read_bytes(), tmp_path)
    save_result(identifier, bank, tmp_path)
    path = job_path(identifier, tmp_path) / "current.json"
    data = json.loads(path.read_text())
    data["result"]["tables"][0]["rows"][0][0] = "tampered"
    path.write_text(json.dumps(data))
    with pytest.raises(ValueError, match="changed"):
        load_result(identifier, tmp_path)


def test_changed_original_source_rejected(bank, tmp_path):
    identifier, source = create_job(
        "dummy.pdf", (tmp_path / "B_bank_errors.pdf").read_bytes(), tmp_path
    )
    save_result(identifier, bank, tmp_path)
    source.write_bytes(b"changed source bytes")
    with pytest.raises(ValueError, match="source changed"):
        load_result(identifier, tmp_path)


@pytest.mark.parametrize("identifier", ["../outside", "C:/Windows", "a" * 31, "a" * 33])
def test_workspace_path_traversal_rejected(tmp_path, identifier):
    with pytest.raises(ValueError):
        job_path(identifier, tmp_path)


def test_verified_checkpoint_avoids_reextraction(tmp_path, monkeypatch):
    invoice_sample(tmp_path)
    source = tmp_path / "A_invoice_register.pdf"
    first = convert(source, checkpoint_dir=tmp_path / "cache")
    import truecell.pipeline as pipeline

    monkeypatch.setattr(
        pipeline,
        "extract_page",
        lambda *args: (_ for _ in ()).throw(AssertionError("must not re-extract")),
    )
    second = convert(source, checkpoint_dir=tmp_path / "cache")
    assert second.tables[0].rows == first.tables[0].rows
    assert all(p.attempts[-1].strategy == "verified-resume" for p in second.pages)


def test_checkpoint_different_signature_recomputed(tmp_path):
    from truecell.models import PageReport

    cache = Checkpoints(tmp_path, "old")
    cache.save(1, [], PageReport(1, "empty"), [])
    assert Checkpoints(tmp_path, "new").load(1) is None
    (tmp_path / "page-001.json").write_text("{bad json")
    assert cache.load(1) is None


def test_export_contains_source_link_and_unreviewed_status(bank, tmp_path):
    from truecell.export import export_workbook

    path = export_workbook(bank, tmp_path / "proof.xlsx")
    workbook = load_workbook(path)
    assert workbook["Data 1"]["F12"].hyperlink.target.startswith("#'Cell Evidence'!")
    assert workbook["Verification"]["B4"].value == "Unreviewed"
    assert workbook["Repair Proposals"].max_row > 1
    assert not any(c.data_type == "f" for s in workbook for row in s for c in row)


def test_exported_cell_evidence_links_original_confidence_and_history(bank, tmp_path):
    from truecell.export import export_workbook

    edited = apply_edit(
        bank,
        0,
        10,
        5,
        "1064.00",
        "Confirmed against the source and corrected the arithmetic value.",
    )
    workbook = load_workbook(export_workbook(edited, tmp_path / "traceable.xlsx"))
    evidence = workbook["Cell Evidence"]
    columns = {cell.value: cell.column for cell in evidence[1]}
    matching = [
        row
        for row in range(2, evidence.max_row + 1)
        if evidence.cell(row, columns["Evidence ID"]).value == "T1:R11:C6"
    ]
    assert matching
    row = matching[0]
    assert evidence.cell(row, columns["Working value"]).value == "1064.00"
    assert evidence.cell(row, columns["Original extracted value"]).value == "1081.00"
    assert evidence.cell(row, columns["Original source reading"]).value == "1081.00"
    assert evidence.cell(row, columns["Confidence type"]).value == "native_text"
    assert evidence.cell(row, columns["Recognition score"]).value is None
    assert (
        evidence.cell(row, columns["Confidence meaning"]).value
        == "Text and geometry came directly from the PDF text layer."
    )
    assert evidence.cell(row, columns["Corrections"]).value == 1
    history = workbook["Correction History"]
    assert history["A2"].value == "T1:R11:C6"
