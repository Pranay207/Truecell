"""HTTP boundaries, offline jobs and persisted source-review workflows."""

import io
import json
import threading
import time
from pathlib import Path

import fitz
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from eval.generate_samples import invoice_sample
from truecell.api import create_app, environment_lan_host
from truecell.checkpoint import source_hash
from truecell.constraints import analyze
from truecell.models import PageReport, Result, Table
from truecell.router import MAX_PAGES
from truecell.workspace import create_job, job_path, set_status


def local_client(app):
    return TestClient(app, base_url="http://127.0.0.1", client=("127.0.0.1", 50000))


def pdf_bytes(pages=1):
    with fitz.open() as document:
        for _ in range(pages):
            document.new_page(width=300, height=400).insert_text((20, 30), "Public dummy document")
        return document.tobytes()


def fake_runner(calls, gate=None):
    def run(source, **options):
        calls.append(options)
        if gate:
            assert gate.wait(10), "Test did not release the worker"
        options["progress"](1, 1)
        rows = [[f"R{i}", str(i), str(i + 10), f"{i * (i + 10):.2f}"] for i in range(1, 7)]
        table = Table(["Reference", "Qty", "Rate", "Amount"], rows, 1, (10, 40, 290, 220), "native")
        table.cell_sources = [
            [
                [
                    {
                        "page": 1,
                        "bbox": [10 + c * 60, 40 + r * 20, 65 + c * 60, 55 + r * 20],
                        "text": value,
                        "engine": "native",
                        "confidence": None,
                    }
                ]
                for c, value in enumerate(row)
            ]
            for r, row in enumerate(rows)
        ]
        result = Result(
            Path(source).name,
            [table],
            pages=[PageReport(1, "digital")],
            source_sha256=source_hash(source),
        )
        analyze(result)
        return result

    return run


def completed(client, identifier):
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        response = client.get(f"/api/jobs/{identifier}")
        assert response.status_code == 200, response.text
        detail = response.json()
        if detail["status"] not in {"ready", "processing"}:
            assert detail["status"] == "complete", detail
            return detail
        time.sleep(0.01)
    pytest.fail("Local worker did not finish")


def upload(client, name="dummy.pdf", content=None, **form):
    return client.post(
        "/api/jobs", files={"file": (name, content or pdf_bytes(), "application/pdf")}, data=form
    )


def test_api_upload_edit_review_source_and_export_workflow(tmp_path):
    calls = []
    app = create_app(tmp_path / "jobs", runner=fake_runner(calls))
    with local_client(app) as client:
        response = upload(client, "../dummy.pdf", verify="true", force_ocr="false")
        assert response.status_code == 202, response.text
        identifier = response.json()["id"]
        detail = completed(client, identifier)
        assert detail["name"] == "dummy.pdf" and detail["progress"] == 1
        assert calls[0]["allow_cloud"] is False and calls[0]["verify"] is True
        assert calls[0]["force_ocr"] is False
        assert client.get("/api/jobs").json()["jobs"][0]["id"] == identifier
        prefix = f"/api/jobs/{identifier}"
        original = detail["result"]["tables"][0]["rows"][0][3]
        evidence = client.get(prefix + "/evidence?table=0&row=0&column=3").json()
        assert evidence["value"] == original and evidence["sources"][0]["bbox"]
        assert client.get(prefix + "/source").content.startswith(b"%PDF")
        assert client.get(prefix + "/pages/1/image").content.startswith(b"\x89PNG")
        assert client.get(prefix + "/pages/1/info").json() == {"width": 300, "height": 400}
        assert client.get(prefix + "/pages/2/image").status_code == 404
        assert client.get(prefix + "/pages/1/image?dpi=900").status_code == 422

        change = {
            "table": 0,
            "row": 0,
            "column": 3,
            "value": "99.00",
            "note": "Checked original",
            "fingerprint": detail["fingerprint"],
        }
        edited_response = client.patch(prefix + "/cells", json=change)
        assert edited_response.status_code == 200, edited_response.text
        edited = edited_response.json()
        assert edited["result"]["tables"][0]["rows"][0][3] == "99.00"
        assert edited["result"]["tables"][0]["cell_sources"][0][3][0]["text"] == original
        assert client.patch(prefix + "/cells", json=change).status_code == 409
        assert client.get(prefix + "/baseline.json").json()["tables"][0]["rows"][0][3] == original
        reviewed = client.post(
            prefix + "/review",
            json={
                "source_checked": True,
                "note": "Compared source and accepted findings",
                "fingerprint": edited["fingerprint"],
            },
        ).json()
        assert reviewed["reviewed"]
        header = {
            "table": 0,
            "row": -1,
            "column": 0,
            "value": "Reference ID",
            "note": "Clarified header",
            "fingerprint": reviewed["fingerprint"],
        }
        changed = client.patch(prefix + "/cells", json=header).json()
        assert not changed["reviewed"] and len(changed["result"]["history"]) == 2
        assert (
            client.post(
                prefix + "/revalidate", json={"fingerprint": changed["fingerprint"]}
            ).status_code
            == 200
        )
        assert len(client.get(prefix + "/audit.json").json()["history"]) == 2
        workbook = load_workbook(io.BytesIO(client.get(prefix + "/export.xlsx").content))
        assert workbook["Data 1"]["D2"].value == 99
        assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
        fork = client.post(prefix + "/process", json={})
        assert fork.status_code == 202 and fork.json()["id"] != identifier


def test_api_priority_review_queue_is_saved_and_source_addressable(tmp_path):
    calls = []

    def uncertain_runner(source, **options):
        result = fake_runner(calls)(source, **options)
        result.tables[0].cell_confidence[(0, 0)] = 0.71
        return result

    app = create_app(tmp_path / "jobs", runner=uncertain_runner)
    with local_client(app) as client:
        identifier = upload(client).json()["id"]
        detail = completed(client, identifier)
        queue = detail["review_queue"]
        item = next(entry for entry in queue["items"] if entry["kind"] == "confidence")
        assert item["status"] == "pending"
        assert (item["table"], item["row"], item["column"], item["page"]) == (0, 0, 0, 1)

        checked = client.post(
            f"/api/jobs/{identifier}/review-items",
            json={
                "item_id": item["id"],
                "source_checked": True,
                "fingerprint": detail["fingerprint"],
            },
        )
        assert checked.status_code == 200, checked.text
        checked_detail = checked.json()
        assert checked_detail["fingerprint"] == detail["fingerprint"]
        assert checked_detail["review_queue"]["pending"] == queue["pending"] - 1
        assert client.get(f"/api/jobs/{identifier}").json()["review_queue"]["checked"] == 1

        invalid = client.post(
            f"/api/jobs/{identifier}/review-items",
            json={"item_id": "missing", "source_checked": True},
        )
        assert invalid.status_code == 400


def test_api_delete_workspace_requires_exact_name_and_preserves_other_jobs(tmp_path):
    root = tmp_path / "jobs"
    app = create_app(root, runner=fake_runner([]))
    with local_client(app) as client:
        first = upload(client, "first.pdf").json()["id"]
        second = upload(client, "second.pdf").json()["id"]
        completed(client, first)
        completed(client, second)

        wrong = client.request(
            "DELETE",
            f"/api/jobs/{first}",
            json={"confirmation": "First.pdf"},
        )
        assert wrong.status_code == 409
        assert job_path(first, root).is_dir()

        deleted = client.request(
            "DELETE",
            f"/api/jobs/{first}",
            json={"confirmation": "first.pdf"},
        )
        assert deleted.status_code == 200
        assert deleted.json() == {"deleted": first, "name": "first.pdf"}
        assert not job_path(first, root).exists()
        assert client.get(f"/api/jobs/{first}").status_code == 404
        assert [record["id"] for record in client.get("/api/jobs").json()["jobs"]] == [second]
        assert client.get(f"/api/jobs/{second}").status_code == 200

        unknown = client.request(
            "DELETE",
            "/api/jobs/" + "f" * 32,
            json={"confirmation": "missing.pdf"},
        )
        assert unknown.status_code == 404
        assert (
            client.request(
                "DELETE",
                "/api/jobs/outside",
                json={"confirmation": "outside"},
            ).status_code
            == 400
        )


def test_api_refuses_to_delete_active_workspace(tmp_path):
    gate = threading.Event()
    root = tmp_path / "jobs"
    app = create_app(root, runner=fake_runner([], gate))
    with local_client(app) as client:
        identifier = upload(client, "active.pdf").json()["id"]
        try:
            response = client.request(
                "DELETE",
                f"/api/jobs/{identifier}",
                json={"confirmation": "active.pdf"},
            )
            assert response.status_code == 409
            assert "finish" in response.json()["detail"]
            assert job_path(identifier, root).is_dir()
        finally:
            gate.set()

        completed(client, identifier)
        response = client.request(
            "DELETE",
            f"/api/jobs/{identifier}",
            json={"confirmation": "active.pdf"},
        )
        assert response.status_code == 200
        assert not job_path(identifier, root).exists()


def test_api_delete_reports_workspace_files_in_use(tmp_path, monkeypatch):
    root = tmp_path / "jobs"
    app = create_app(root, runner=fake_runner([]))
    with local_client(app) as client:
        identifier = upload(client, "locked.pdf").json()["id"]
        completed(client, identifier)

        def locked(_folder):
            raise PermissionError("file is open")

        with monkeypatch.context() as patch:
            patch.setattr("truecell.api.shutil.rmtree", locked)
            response = client.request(
                "DELETE",
                f"/api/jobs/{identifier}",
                json={"confirmation": "locked.pdf"},
            )

        assert response.status_code == 409
        assert "in use" in response.json()["detail"]
        assert job_path(identifier, root).is_dir()


def test_api_limits_invalid_input_and_addresses(tmp_path):
    app = create_app(tmp_path / "jobs", runner=fake_runner([]), max_upload_bytes=100_000)
    with local_client(app) as client:
        assert upload(client, content=b"x" * 100_001).status_code == 413
        assert upload(client, "script.exe").status_code == 400
        assert upload(client, content=b"not a PDF").status_code == 400
        long_document = upload(client, content=pdf_bytes(116))
        assert long_document.status_code == 202
        completed(client, long_document.json()["id"])
        assert upload(client, content=pdf_bytes(MAX_PAGES + 1)).status_code == 400
        assert client.get("/api/jobs/invalid").status_code == 400
        assert client.get("/api/jobs/" + "a" * 32).status_code == 404
        response = upload(client)
        identifier = response.json()["id"]
        completed(client, identifier)
        prefix = f"/api/jobs/{identifier}"
        for table, row, column in [(-1, 0, 0), (0, -2, 0), (0, 0, -1), (0, 999, 0)]:
            change = {"table": table, "row": row, "column": column, "value": "x", "note": "checked"}
            assert client.patch(prefix + "/cells", json=change).status_code == 400
        assert (
            client.patch(
                prefix + "/cells",
                json={"table": False, "row": 0, "column": 0, "value": "x", "note": "checked"},
            ).status_code
            == 422
        )
        assert client.post(prefix + "/process", json={"allow_cloud": True}).status_code == 422
        assert (
            client.post(
                prefix + "/review", json={"source_checked": False, "note": "checked"}
            ).status_code
            == 400
        )


def test_api_blocks_remote_origins_and_dns_rebinding(tmp_path):
    app = create_app(tmp_path / "jobs", runner=fake_runner([]))
    with local_client(app) as client:
        health = client.get("/api/health").json()
        assert health["offline"] is True
        assert health["max_pages"] == 200 == MAX_PAGES
        assert (
            client.get("/api/health", headers={"Origin": "https://evil.example"}).status_code == 403
        )
        assert client.get("/api/health", headers={"Origin": "http://["}).status_code == 403
        assert client.get("/api/health", headers={"Host": "evil.example"}).status_code == 403
        assert (
            client.get("/api/health", headers={"Origin": "http://localhost:5173"}).status_code
            == 200
        )
        preflight = client.options(
            "/api/jobs/" + "a" * 32,
            headers={
                "Origin": "http://localhost:5173",
                "Access-Control-Request-Method": "DELETE",
            },
        )
        assert preflight.status_code == 200
        assert "DELETE" in preflight.headers["access-control-allow-methods"]


def test_api_rejects_remote_transport(tmp_path):
    remote_app = create_app(tmp_path / "remote-jobs", runner=fake_runner([]))
    with TestClient(
        remote_app, base_url="http://127.0.0.1", client=("198.51.100.2", 50000)
    ) as remote:
        assert remote.get("/api/health").status_code == 403


def test_api_lan_mode_allows_exact_private_host_and_same_origin(tmp_path):
    app = create_app(tmp_path / "lan-jobs", runner=fake_runner([]), lan_host="192.168.10.20")
    with TestClient(
        app,
        base_url="http://192.168.10.20:8501",
        client=("192.168.10.50", 50000),
    ) as phone:
        health = phone.get("/api/health")
        assert health.status_code == 200
        assert health.json()["network_mode"] == "lan"
        assert (
            phone.get("/api/health", headers={"Origin": "http://192.168.10.20:8501"}).status_code
            == 200
        )
        assert (
            phone.get("/api/health", headers={"Origin": "http://192.168.10.20:8502"}).status_code
            == 403
        )
        assert phone.get("/api/health", headers={"Host": "192.168.10.21:8501"}).status_code == 403

    public_client = TestClient(
        create_app(tmp_path / "public-jobs", runner=fake_runner([]), lan_host="192.168.10.20"),
        base_url="http://192.168.10.20:8501",
        client=("8.8.8.8", 50000),
    )
    with public_client:
        assert public_client.get("/api/health").status_code == 403


def test_api_lan_configuration_fails_closed(monkeypatch, tmp_path):
    with pytest.raises(ValueError, match="RFC 1918"):
        create_app(tmp_path / "invalid-lan", runner=fake_runner([]), lan_host="198.51.100.4")
    monkeypatch.setenv("TRUECELL_LAN_HOST", "evil.example")
    with pytest.raises(RuntimeError, match="RFC 1918"):
        environment_lan_host()


def test_api_source_and_saved_payload_tampering_is_rejected(tmp_path):
    root = tmp_path / "jobs"
    with local_client(create_app(root, runner=fake_runner([]))) as client:
        identifier = upload(client).json()["id"]
        completed(client, identifier)
        prefix = f"/api/jobs/{identifier}"
        folder = job_path(identifier, root)
        source = folder / "source.pdf"
        original = source.read_bytes()
        source.write_bytes(b"changed source")
        assert client.get(prefix).status_code == 409
        assert client.get(prefix + "/source").status_code == 409
        assert client.get(prefix + "/pages/1/image").status_code == 409
        source.write_bytes(original)
        current = folder / "current.json"
        payload = json.loads(current.read_text("utf-8"))
        payload["result"]["tables"][0]["rows"][0][0] = "tampered"
        current.write_text(json.dumps(payload), "utf-8")
        assert client.get(prefix + "/audit.json").status_code == 409
        assert client.get(prefix + "/baseline.json").status_code == 200


def test_api_rejects_source_path_escape_in_saved_metadata(tmp_path):
    root = tmp_path / "jobs"
    with local_client(create_app(root, runner=fake_runner([]))) as client:
        identifier = upload(client).json()["id"]
        completed(client, identifier)
        metadata = job_path(identifier, root) / "metadata.json"
        record = json.loads(metadata.read_text("utf-8"))
        record["source"] = "../outside.pdf"
        metadata.write_text(json.dumps(record), "utf-8")
        assert client.get(f"/api/jobs/{identifier}/source").status_code == 409
        assert client.get(f"/api/jobs/{identifier}/pages/1/image").status_code == 409


def test_api_worker_queue_is_bounded_and_edits_wait_for_processing(tmp_path):
    gate = threading.Event()
    calls = []
    root = tmp_path / "jobs"
    with local_client(create_app(root, runner=fake_runner(calls, gate))) as client:
        try:
            responses = [upload(client, f"dummy{i}.pdf") for i in range(3)]
            assert all(response.status_code == 202 for response in responses)
            assert upload(client).status_code == 429
            assert len(client.get("/api/jobs").json()["jobs"]) == 3
            identifier = responses[0].json()["id"]
            prefix = f"/api/jobs/{identifier}"
            assert client.post(prefix + "/process", json={}).status_code == 409
            assert (
                client.patch(
                    prefix + "/cells",
                    json={"table": 0, "row": 0, "column": 0, "value": "x", "note": "checked"},
                ).status_code
                == 409
            )
            assert len(calls) == 1
        finally:
            gate.set()
        for response in responses:
            completed(client, response.json()["id"])


def test_api_samples_restart_recovery_and_missing_frontend(tmp_path):
    samples = tmp_path / "samples"
    samples.mkdir()
    (samples / "dummy.pdf").write_bytes(pdf_bytes())
    root = tmp_path / "jobs"
    identifier, _ = create_job("interrupted.pdf", pdf_bytes(), root)
    set_status(identifier, "processing", root)
    app = create_app(root, samples=samples, runner=fake_runner([]), frontend=tmp_path / "missing")
    with local_client(app) as client:
        assert client.get("/").status_code == 503
        assert client.get("/api/health").status_code == 200
        assert client.get(f"/api/jobs/{identifier}").json()["status"] == "failed"
        assert client.post(f"/api/jobs/{identifier}/process", json={}).status_code == 202
        completed(client, identifier)
        assert client.get("/api/samples").json()["samples"][0]["filename"] == "dummy.pdf"
        response = client.post("/api/samples/dummy.pdf")
        assert response.status_code == 202
        completed(client, response.json()["id"])
        assert client.post("/api/samples/not-allowed.pdf").status_code == 404


def test_api_serves_prebuilt_frontend_without_node(tmp_path):
    frontend = tmp_path / "dist"
    frontend.mkdir()
    (frontend / "index.html").write_text("<html>Truecell Studio</html>", "utf-8")
    with local_client(
        create_app(tmp_path / "jobs", frontend=frontend, runner=fake_runner([]))
    ) as client:
        assert "Truecell Studio" in client.get("/").text
        assert client.get("/api/health").status_code == 200


def test_api_runs_real_native_converter_and_preserves_all_rows(tmp_path):
    invoice_sample(tmp_path)
    source = tmp_path / "A_invoice_register.pdf"
    with local_client(create_app(tmp_path / "jobs")) as client:
        response = upload(client, source.name, source.read_bytes())
        assert response.status_code == 202
        detail = completed(client, response.json()["id"])
        result = detail["result"]
        assert len(result["tables"]) == 1 and len(result["tables"][0]["rows"]) == 48
        assert len(result["tables"][0]["headers"]) == 12
        assert detail["coverage"]["cells_with_source_boxes"] == detail["coverage"]["cells"]


def test_api_failed_worker_releases_queue_and_partial_results_remain_available(tmp_path):
    def worker(source, **options):
        if (
            "failed"
            in json.loads((Path(source).parent / "metadata.json").read_text("utf-8"))["name"]
        ):
            raise RuntimeError("Dummy extraction failure")
        result = fake_runner([])(source, **options)
        result.pages[0].error = "Dummy partial page finding"
        return result

    with local_client(create_app(tmp_path / "jobs", runner=worker)) as client:
        first = upload(client, "failed.pdf").json()["id"]
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            failed = client.get(f"/api/jobs/{first}").json()
            if failed["status"] == "failed":
                break
            time.sleep(0.01)
        assert failed["status"] == "failed" and "Dummy extraction failure" in failed["error"]
        second = upload(client, "partial.pdf").json()["id"]
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            partial = client.get(f"/api/jobs/{second}").json()
            if partial["status"] == "partial":
                break
            time.sleep(0.01)
        assert partial["status"] == "partial" and partial["result"]["tables"]
        assert client.get(f"/api/jobs/{second}/audit.json").status_code == 200
        resumed = client.post(f"/api/jobs/{second}/process", json={})
        assert resumed.status_code == 202 and resumed.json()["id"] == second


def test_api_preview_and_challenge_do_not_change_saved_snapshot_or_history(tmp_path):
    root = tmp_path / "jobs"
    with local_client(create_app(root, runner=fake_runner([]))) as client:
        identifier = upload(client).json()["id"]
        detail = completed(client, identifier)
        prefix = f"/api/jobs/{identifier}"
        broken = client.patch(
            prefix + "/cells",
            json={
                "table": 0,
                "row": 0,
                "column": 3,
                "value": "99.00",
                "note": "Deliberate test edit",
                "fingerprint": detail["fingerprint"],
            },
        ).json()
        before = {
            path.name: path.read_bytes()
            for path in job_path(identifier, root).iterdir()
            if path.is_file()
        }
        response = client.post(
            prefix + "/preview",
            json={
                "table": 0,
                "row": 0,
                "column": 3,
                "value": "11.00",
                "fingerprint": broken["fingerprint"],
            },
        )
        assert response.status_code == 200, response.text
        preview = response.json()
        assert preview["saved"] is False and preview["proposed"] == "11.00"
        assert preview["observed"] == "99.00"
        assert (
            preview["before"]["failed_equation_checks"] > preview["after"]["failed_equation_checks"]
        )
        report_response = client.post(
            prefix + "/challenge",
            json={
                "table": 0,
                "fingerprint": broken["fingerprint"],
            },
        )
        assert report_response.status_code == 200, report_response.text
        report = report_response.json()
        assert report["saved"] is False and 0 < report["tested"] <= 24
        assert len(report["cases"]) == report["tested"]
        assert report["rows_tested"] == 6 and report["row_limit"] == 120
        assert "does not test OCR" in report["scope"]
        after = {
            path.name: path.read_bytes()
            for path in job_path(identifier, root).iterdir()
            if path.is_file()
        }
        assert after == before
        unchanged = client.get(prefix).json()
        assert unchanged["fingerprint"] == broken["fingerprint"]
        assert unchanged["result"] == broken["result"]
        assert len(unchanged["result"]["history"]) == 1


def test_api_preview_and_challenge_reject_stale_or_invalid_requests(tmp_path):
    with local_client(create_app(tmp_path / "jobs", runner=fake_runner([]))) as client:
        identifier = upload(client).json()["id"]
        detail = completed(client, identifier)
        prefix = f"/api/jobs/{identifier}"
        for table, row, column in [(-1, 0, 0), (0, -2, 0), (0, 0, -1), (0, 999, 0)]:
            assert (
                client.post(
                    prefix + "/preview",
                    json={
                        "table": table,
                        "row": row,
                        "column": column,
                        "value": "x",
                        "fingerprint": detail["fingerprint"],
                    },
                ).status_code
                == 400
            )
        assert (
            client.post(
                prefix + "/preview",
                json={
                    "table": 0,
                    "row": 0,
                    "column": 3,
                    "value": "99.00",
                    "fingerprint": "stale",
                },
            ).status_code
            == 409
        )
        assert (
            client.post(
                prefix + "/challenge", json={"table": 0, "fingerprint": "stale"}
            ).status_code
            == 409
        )
        assert client.post(prefix + "/challenge", json={"table": -1}).status_code == 400
        assert client.post(prefix + "/challenge", json={"table": False}).status_code == 422


def test_api_challenge_caps_large_tables_without_dropping_saved_rows(tmp_path):
    def larger_table(source, **options):
        result = fake_runner([])(source, **options)
        table = result.tables[0]
        for index in range(7, 131):
            table.rows.append([f"R{index}", str(index), str(index + 10), str(index * (index + 10))])
            table.provenance.append([(1, index)])
            table.cell_sources.append([[] for _ in table.headers])
            table.row_engines.append(table.engine)
        analyze(result)
        return result

    with local_client(create_app(tmp_path / "jobs", runner=larger_table)) as client:
        identifier = upload(client).json()["id"]
        detail = completed(client, identifier)
        response = client.post(
            f"/api/jobs/{identifier}/challenge",
            json={
                "table": 0,
                "fingerprint": detail["fingerprint"],
            },
        )
        assert response.status_code == 200, response.text
        report = response.json()
        assert report["rows_available"] == 130 and report["rows_tested"] == 120
        assert report["tested"] == 24 and all(case["row"] <= 120 for case in report["cases"])
        unchanged = client.get(f"/api/jobs/{identifier}").json()
        assert unchanged["fingerprint"] == detail["fingerprint"]
        assert len(unchanged["result"]["tables"][0]["rows"]) == 130


@pytest.mark.parametrize("edited", [False, True])
def test_api_reprocessing_preserves_reviewed_workspace_and_returns_fresh_job(tmp_path, edited):
    root = tmp_path / "jobs"
    calls = []
    with local_client(create_app(root, runner=fake_runner(calls))) as client:
        identifier = upload(client, verify="true", force_ocr="true").json()["id"]
        detail = completed(client, identifier)
        prefix = f"/api/jobs/{identifier}"
        if edited:
            detail = client.patch(
                prefix + "/cells",
                json={
                    "table": 0,
                    "row": 0,
                    "column": 3,
                    "value": "99.00",
                    "note": "Checked source",
                    "fingerprint": detail["fingerprint"],
                },
            ).json()
        else:
            detail = client.post(
                prefix + "/review",
                json={
                    "source_checked": True,
                    "note": "Compared original pages",
                    "fingerprint": detail["fingerprint"],
                },
            ).json()
        original_files = {
            path.name: path.read_bytes()
            for path in job_path(identifier, root).iterdir()
            if path.is_file()
        }
        response = client.post(prefix + "/process")
        assert response.status_code == 202, response.text
        new_id = response.json()["id"]
        assert new_id != identifier and response.json()["forked_from"] == identifier
        fresh = completed(client, new_id)
        assert fresh["name"] == detail["name"]
        assert fresh["result"]["source_sha256"] == detail["result"]["source_sha256"]
        assert fresh["result"]["history"] == [] and not fresh["reviewed"]
        assert fresh["result"]["tables"][0]["rows"][0][3] == "11.00"
        assert calls[-1]["verify"] is True and calls[-1]["force_ocr"] is True
        assert client.get(prefix).json()["result"] == detail["result"]
        assert original_files == {
            path.name: path.read_bytes()
            for path in job_path(identifier, root).iterdir()
            if path.is_file()
        }
